# =====================================================================
# Project: paint-crosschecker
# Copyright (c) 2026 kmm851010-maker. All rights reserved.
# Unauthorized copying, modification, or distribution is strictly prohibited.
# =====================================================================
"""
생산계획서 복합 표(17열) 초정밀 추출 모듈
Anthropic Tool Use(tool_choice) 강제 적용으로 JSON 스키마 100% 보장.
2단계: AI가 표 전체를 읽기 → 코드로 신규 필터링.
"""

import json
import re

import anthropic

from utils.helpers import (
    encode_image_to_base64,
    detect_media_type,
    parse_quantity_text,
    normalize_color_code,
    is_valid_item_code,
    auto_correct_code,
    is_document_file,
)


# Tool Use 스키마 정의
EXTRACT_TABLE_TOOL = {
    "name": "extract_table_data",
    "description": "생산계획표 이미지에서 표 데이터를 추출합니다. 모든 행과 열을 빠짐없이 읽어 반환합니다.",
    "input_schema": {
        "type": "object",
        "properties": {
            "headers": {
                "type": "array",
                "items": {"type": "string"},
                "description": "표의 헤더(컬럼명) 리스트. 왼쪽부터 순서대로.",
            },
            "rows": {
                "type": "array",
                "items": {
                    "type": "array",
                    "items": {
                        "oneOf": [
                            {"type": "string"},
                            {"type": "number"},
                            {"type": "null"},
                        ]
                    },
                },
                "description": "데이터 행 리스트. 각 행은 셀 값의 배열.",
            },
        },
        "required": ["headers", "rows"],
    },
}

TABLE_READ_PROMPT = """이 이미지의 표를 **모든 셀을 빠짐없이** 정확히 읽으세요.

규칙:
- 표의 모든 행과 모든 열을 왼쪽→오른쪽, 위→아래 순서로 읽으세요.
- 빈 셀은 null로 표시하세요.
- 숫자는 숫자 타입으로, 텍스트는 문자열로 반환하세요.
- "16(수)", "3(화)", "22(목)" 같은 숫자+괄호 형식은 문자열 그대로 반환하세요. 숫자로 변환하지 마세요!
- "창고3", "13중12", "위생산" 등 텍스트 포함 셀도 문자열 그대로 반환하세요.
- 품목코드(영문+숫자 7자리)는 한 글자도 틀리지 않게 정확히 읽으세요.
- 0↔O, 1↔I, 5↔S, 8↔B 혼동 주의.
- "재고", "신규" 등 헤더도 정확히 읽으세요.

extract_table_data 도구를 사용하여 결과를 반환하세요."""


def extract_table_from_image_tool_use(
    image_bytes: bytes,
    file_name: str,
    api_key: str,
    model: str = "claude-opus-4-8",
) -> dict:
    """
    Tool Use 강제로 이미지에서 표 데이터를 추출합니다.
    Returns: {"headers": [...], "rows": [[...], ...]}
    """
    client = anthropic.Anthropic(api_key=api_key)
    b64 = encode_image_to_base64(image_bytes)
    media_type = detect_media_type(file_name)

    response = client.messages.create(
        model=model,
        max_tokens=16000,
        tools=[EXTRACT_TABLE_TOOL],
        tool_choice={"type": "tool", "name": "extract_table_data"},
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                {"type": "text", "text": TABLE_READ_PROMPT},
            ],
        }],
    )

    # Tool Use 응답에서 입력 추출
    for block in response.content:
        if getattr(block, "type", "") == "tool_use" and block.name == "extract_table_data":
            return block.input

    raise ValueError("Tool Use 응답을 받지 못했습니다.")


def _find_new_columns(headers: list) -> list:
    """헤더에서 '신규' 열 인덱스를 찾습니다."""
    new_cols = []
    for i, h in enumerate(headers):
        if h and "신규" in str(h):
            new_cols.append(i)
    return new_cols


def _find_stock_columns(headers: list) -> list:
    """헤더에서 '재고' 열 인덱스를 찾습니다."""
    stock_cols = []
    for i, h in enumerate(headers):
        if h and "재고" in str(h):
            stock_cols.append(i)
    return stock_cols


def extract_new_items_from_table(table_data: dict) -> list:
    """
    추출된 표에서 프로그래밍 로직으로 '신규' 품목만 필터링합니다.
    AI 판단 없이 코드로 정확하게 처리.
    """
    headers = table_data.get("headers", [])
    rows = table_data.get("rows", [])

    if not headers or not rows:
        return []

    new_cols = _find_new_columns(headers)
    stock_cols = _find_stock_columns(headers)

    if not new_cols:
        # '신규' 헤더를 못 찾으면 4열 반복 패턴으로 추정
        total = len(headers)
        data_cols = total - 1 if total % 4 == 1 else total
        new_cols = [i for i in range(3, data_cols, 4)]

    items = []
    for row in rows:
        for new_col in new_cols:
            if new_col >= len(row):
                continue

            # 신규 셀 값
            cell_val = row[new_col]
            parsed = parse_quantity_text(cell_val)
            qty = parsed["quantity"]

            if qty <= 0:
                continue

            # 신규 셀이 재고 열과 같은 위치면 스킵
            if new_col in stock_cols:
                continue

            # 역추적: 왼쪽으로 품목코드와 제조사 찾기
            code = ""
            maker = ""

            # 품목코드: 신규 기준 왼쪽 3칸 (같은 블록의 첫 열)
            code_col = new_col - 3
            if 0 <= code_col < len(row) and row[code_col]:
                raw_code = str(row[code_col]).strip()
                corrected = auto_correct_code(raw_code)
                if is_valid_item_code(corrected):
                    code = corrected

            # 품목코드를 못 찾으면 왼쪽으로 탐색
            if not code:
                for offset in range(1, min(4, new_col + 1)):
                    check_col = new_col - offset
                    if 0 <= check_col < len(row) and row[check_col]:
                        raw = str(row[check_col]).strip()
                        corrected = auto_correct_code(raw)
                        if is_valid_item_code(corrected):
                            code = corrected
                            # 제조사는 코드 바로 오른쪽
                            maker_col = check_col + 1
                            if maker_col < len(row) and row[maker_col] and maker_col != new_col:
                                maker = str(row[maker_col]).strip()
                            break

            # 제조사를 아직 못 찾았으면 코드+1 위치
            if not maker and code:
                maker_col = new_col - 2
                if 0 <= maker_col < len(row) and row[maker_col]:
                    m = str(row[maker_col]).strip()
                    if not is_valid_item_code(auto_correct_code(m)):
                        maker = m

            if not code:
                continue

            # "위생산" 오인식 제거
            if "위생산" in maker:
                maker = ""
            if "위생산" in code:
                continue  # 코드 자체가 위생산이면 유효한 품목이 아님

            items.append({
                "색상코드": code,
                "제조사": maker,
                "신규": qty,
                "비고": parsed.get("schedule_day", ""),
                "라인": "",
                "위치": "",
                "재고": 0,
                "생산량": 0,
            })

    # 같은 품목코드 합산
    merged = {}
    for item in items:
        code = item["색상코드"]
        if code in merged:
            merged[code]["신규"] += item["신규"]
            if item["비고"] and not merged[code]["비고"]:
                merged[code]["비고"] = item["비고"]
        else:
            merged[code] = item.copy()

    return list(merged.values())


def _excel_to_table_data(file_bytes: bytes, file_name: str) -> dict:
    """Excel/CSV 파일을 table_data 형식으로 변환."""
    from io import BytesIO
    import pandas as pd
    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    is_ole = len(file_bytes) >= 8 and file_bytes[:8] == bytes.fromhex("d0cf11e0a1b011ae")
    try:
        if ext == "csv":
            for enc in ["utf-8", "cp949", "euc-kr", "latin-1"]:
                try:
                    df = pd.read_csv(BytesIO(file_bytes), encoding=enc)
                    break
                except Exception:
                    continue
            else:
                return {"headers": [], "rows": []}
        elif is_ole or ext == "xls":
            df = pd.read_excel(BytesIO(file_bytes), engine="xlrd")
        else:
            df = pd.read_excel(BytesIO(file_bytes), engine="openpyxl")
        headers = [str(c) for c in df.columns.tolist()]
        rows = [[str(v) if v != "" else "" for v in row] for row in df.fillna("").values.tolist()]
        return {"headers": headers, "rows": rows}
    except Exception:
        return {"headers": [], "rows": []}


def _extract_table_from_pdf(pdf_bytes: bytes, api_key: str, model: str) -> dict:
    """PDF에서 표 데이터를 추출합니다 (Anthropic native PDF support)."""
    import base64 as _b64
    client = anthropic.Anthropic(api_key=api_key)
    b64 = _b64.standard_b64encode(pdf_bytes).decode("utf-8")
    response = client.messages.create(
        model=model,
        max_tokens=16000,
        tools=[EXTRACT_TABLE_TOOL],
        tool_choice={"type": "tool", "name": "extract_table_data"},
        messages=[{
            "role": "user",
            "content": [
                {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": b64}},
                {"type": "text", "text": TABLE_READ_PROMPT},
            ],
        }],
    )
    for block in response.content:
        if getattr(block, "type", "") == "tool_use" and block.name == "extract_table_data":
            return block.input
    raise ValueError("PDF Tool Use 응답을 받지 못했습니다.")


def _extract_table_from_word(word_bytes: bytes) -> dict:
    """Word(.docx) 파일에서 표 데이터를 추출합니다."""
    from io import BytesIO
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx가 필요합니다. pip install python-docx")
    doc = docx.Document(BytesIO(word_bytes))
    if not doc.tables:
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return {"headers": ["내용"], "rows": [[line] for line in lines]}
    best_table = max(doc.tables, key=lambda t: len(t.rows) * len(t.columns))
    all_rows = [[cell.text.strip() for cell in row.cells] for row in best_table.rows]
    if not all_rows:
        return {"headers": [], "rows": []}
    return {"headers": all_rows[0], "rows": all_rows[1:]}


def extract_production_plan(
    file_bytes: bytes,
    file_name: str,
    api_key: str,
    model: str = "claude-opus-4-8",
) -> dict:
    """
    생산계획서(이미지/엑셀/PDF/Word)에서 신규 입고 대상 품목을 추출합니다.

    Returns: {
        "items": [{"색상코드", "제조사", "신규", ...}, ...],
        "table_data": {"headers": [...], "rows": [[...], ...]} or None,
    }
    """
    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""

    # PDF → Claude native document support
    if ext == "pdf":
        table_data = _extract_table_from_pdf(file_bytes, api_key, model)
        items = extract_new_items_from_table(table_data)
        return {"items": items, "table_data": table_data}

    # Word → python-docx 파싱
    if ext in ("docx", "doc"):
        table_data = _extract_table_from_word(file_bytes)
        items = extract_new_items_from_table(table_data)
        return {"items": items, "table_data": table_data}

    # Excel/CSV → 기존 파서 + table_data 합성
    if is_document_file(file_name, file_bytes):
        from modules.plan_excel_parser import parse_plan_excel
        items = parse_plan_excel(file_bytes, file_name)
        table_data = _excel_to_table_data(file_bytes, file_name)
        return {"items": items, "table_data": table_data}

    # 이미지 → table_extractor로 표 전체 읽기
    from modules.table_extractor import extract_table_from_image
    table_data = extract_table_from_image(file_bytes, file_name, api_key, model)
    items = extract_new_items_from_table(table_data)
    return {"items": items, "table_data": table_data}


# 하위 호환용
def extract_document_data(file_bytes, file_name, api_key, model="claude-opus-4-8"):
    """기존 호환: 범용 문서 데이터 추출 (ERP 파서 등에서 사용)"""
    if is_document_file(file_name, file_bytes):
        return _parse_document_to_universal(file_bytes, file_name)

    # ERP 이미지: 기존 프롬프트 방식 유지
    client = anthropic.Anthropic(api_key=api_key)
    b64 = encode_image_to_base64(file_bytes)
    media_type = detect_media_type(file_name)

    ERP_PROMPT = """이 이미지에서 표 데이터를 정확히 추출하세요.
품목코드는 영문+숫자 정확히. 숫자는 숫자형으로.
JSON만 응답: {"doc_type":"receipt","items":[{"color_code":"코드","quantity":숫자,"quantity_label":"입고","weight_kg":0,"manufacturer":"제조사","extra_info":""}]}"""

    response = client.messages.create(
        model=model,
        max_tokens=16000,
        messages=[{
            "role": "user",
            "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                {"type": "text", "text": ERP_PROMPT},
            ],
        }],
    )

    text = ""
    for block in response.content:
        if getattr(block, "type", "") == "text":
            text = getattr(block, "text", "")
            if text:
                break

    if not text:
        raise ValueError("OCR 응답에서 텍스트를 찾을 수 없습니다.")

    return _parse_json_response(text)


def _parse_json_response(text: str) -> dict:
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    match = re.search(r"```(?:json)?\s*([\s\S]*?)```", text)
    if match:
        raw = match.group(1).strip()
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            fixed = re.sub(r",\s*([}\]])", r"\1", raw)
            try:
                return json.loads(fixed)
            except json.JSONDecodeError:
                pass
    match = re.search(r"\{[\s\S]*\}", text)
    if match:
        raw = match.group(0)
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            fixed = re.sub(r",\s*([}\]])", r"\1", raw)
            try:
                return json.loads(fixed)
            except json.JSONDecodeError:
                pass
    raise ValueError(f"JSON 파싱 실패: {text[:300]}")


def _parse_document_to_universal(file_bytes: bytes, file_name: str) -> dict:
    from io import BytesIO
    import pandas as pd

    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    is_ole = len(file_bytes) >= 8 and file_bytes[:8] == bytes.fromhex("d0cf11e0a1b011ae")

    if ext == "csv":
        for enc in ["utf-8", "cp949", "euc-kr", "latin-1"]:
            try:
                df = pd.read_csv(BytesIO(file_bytes), encoding=enc)
                break
            except Exception:
                continue
        else:
            raise ValueError("CSV 인코딩 인식 불가")
    elif is_ole or ext == "xls":
        df = pd.read_excel(BytesIO(file_bytes), engine="xlrd")
    else:
        df = pd.read_excel(BytesIO(file_bytes), engine="openpyxl")

    color_kw = ["색상", "품목", "코드", "color", "품명", "clrcd"]
    qty_kw = ["신규", "수량", "new", "qty", "입고", "drum"]
    weight_kw = ["중량", "무게", "kg", "weight"]
    mfr_kw = ["제조", "maker", "회사"]

    color_col = qty_col = weight_col = mfr_col = None
    for col in df.columns:
        cl = str(col).lower()
        if not color_col and any(k in cl for k in color_kw): color_col = col
        if not qty_col and any(k in cl for k in qty_kw): qty_col = col
        if not weight_col and any(k in cl for k in weight_kw): weight_col = col
        if not mfr_col and any(k in cl for k in mfr_kw): mfr_col = col

    if not color_col and df.columns.tolist():
        color_col = df.columns[0]

    items = []
    for _, row in df.iterrows():
        code = str(row[color_col]).strip() if color_col and pd.notna(row[color_col]) else ""
        if not code or code == "nan":
            continue
        qty = 0
        if qty_col and pd.notna(row[qty_col]):
            try:
                qty = int(float(row[qty_col]))
            except (ValueError, TypeError):
                qty = 1
        wgt = 0
        if weight_col and pd.notna(row[weight_col]):
            try:
                wgt = float(row[weight_col])
            except (ValueError, TypeError):
                pass
        mfr = str(row[mfr_col]).strip() if mfr_col and pd.notna(row[mfr_col]) else ""

        items.append({
            "color_code": code,
            "quantity": qty if qty > 0 else 1,
            "quantity_label": str(qty_col) if qty_col else "행",
            "weight_kg": wgt,
            "manufacturer": mfr,
            "extra_info": "",
        })

    doc_type = "unknown"
    if qty_col:
        ql = str(qty_col).lower()
        if any(k in ql for k in ["신규", "new"]): doc_type = "plan"
        elif any(k in ql for k in ["입고", "drum"]): doc_type = "receipt"

    return {"doc_type": doc_type, "items": items}


def extract_erp_from_image(file_bytes, file_name, api_key, model="claude-opus-4-8"):
    return extract_document_data(file_bytes, file_name, api_key, model)


EXTRACT_LOT_TOOL = {
    "name": "extract_lot_list",
    "description": "재고 문서에서 LOT번호와 품명 목록을 추출합니다.",
    "input_schema": {
        "type": "object",
        "properties": {
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "lot": {"type": "string", "description": "LOT번호 (9자리: 영문1 + 숫자8, 예: G12345678)"},
                        "product": {"type": "string", "description": "품명 또는 색상코드"},
                    },
                    "required": ["lot", "product"],
                },
                "description": "LOT번호와 품명 쌍 목록",
            }
        },
        "required": ["items"],
    },
}

LOT_EXTRACT_PROMPT = """이 문서에서 LOT번호와 품명을 모두 추출하세요.

LOT번호: 영문 대문자 1글자 + 숫자 8자리 = 총 9자리 (예: G12345678, D87654321)
품명: LOT번호와 연결된 제품명 또는 색상코드

extract_lot_list 도구를 사용하여 모든 항목을 반환하세요."""


def _extract_lots_from_raw_text(file_bytes: bytes) -> list:
    """파일 바이트에서 LOT 패턴을 직접 정규식 스캔 (포맷 불명 파일 폴백)."""
    import re
    maker_codes = {"G", "D", "K", "S", "Y", "P"}
    lot_pat = re.compile(r"[A-Z][A-Z0-9]{7,11}")
    result = []
    seen = set()
    for enc in ["utf-8", "cp949", "euc-kr", "latin-1", "utf-16"]:
        try:
            text = file_bytes.decode(enc, errors="replace")
            for m in lot_pat.finditer(text):
                lot = m.group()
                if lot[0] not in maker_codes:
                    continue
                if len(lot) == 10:
                    lot = lot[:-1]
                if lot in seen:
                    continue
                seen.add(lot)
                result.append({"lot": lot, "product": ""})
            if result:
                return result
        except Exception:
            continue
    return result


def _extract_lots_from_excel(file_bytes: bytes, filename: str) -> list:
    """엑셀/CSV에서 LOT번호+품명 추출 (Vision AI 없이 직접 파싱)."""
    import re
    from io import BytesIO
    import pandas as pd

    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    is_ole = len(file_bytes) >= 8 and file_bytes[:8] == bytes.fromhex("d0cf11e0a1b011ae")

    if ext == "csv":
        for enc in ["utf-8", "cp949", "euc-kr", "latin-1"]:
            try:
                df = pd.read_csv(BytesIO(file_bytes), encoding=enc, dtype=str)
                break
            except Exception:
                continue
        else:
            raise ValueError("CSV 인코딩 인식 불가")
    elif is_ole or ext == "xls":
        df = pd.read_excel(BytesIO(file_bytes), engine="xlrd", dtype=str)
    else:
        try:
            df = pd.read_excel(BytesIO(file_bytes), engine="openpyxl", dtype=str)
        except Exception:
            try:
                # xlsx 확장자지만 실제 XLS 포맷인 경우 xlrd로 폴백
                df = pd.read_excel(BytesIO(file_bytes), engine="xlrd", dtype=str)
            except Exception:
                # 엑셀 파싱 완전 실패 → 텍스트에서 LOT 패턴 직접 스캔
                return _extract_lots_from_raw_text(file_bytes)

    df = df.fillna("")

    # LOT 패턴: 영문1+숫자8 (바코드형) 또는 영문1+영숫자8~11 (재고장형)
    # 예) G12345678, D26A304031, D26C43601
    lot_pat = re.compile(r"^[A-Z][A-Z0-9]{7,11}$")
    # 제조사 코드 확인 (첫 글자)
    maker_codes = {"G", "D", "K", "S", "Y", "P"}

    lot_kw = ["lot", "로트", "lot-no", "lot번호", "lot no"]
    prod_kw = ["품명", "제품명", "product", "색상", "품목"]

    cols = list(df.columns)
    lot_col = prod_col = None
    for col in cols:
        cl = str(col).lower()
        if not lot_col and any(k in cl for k in lot_kw):
            lot_col = col
        if not prod_col and any(k in cl for k in prod_kw):
            prod_col = col

    # 헤더에서 못 찾으면 데이터 내 LOT 패턴 열 자동 감지
    if not lot_col:
        for col in cols:
            sample = df[col].dropna().head(30)
            matched = sample.apply(lambda v: bool(lot_pat.match(str(v).strip().upper())) and str(v).strip().upper()[:1] in maker_codes).sum()
            if matched >= 3:
                lot_col = col
                break

    if not lot_col:
        raise ValueError("LOT번호 열을 찾을 수 없습니다. (열 이름에 'LOT' 또는 '로트' 포함 필요)")

    # 품명 열을 못 찾았으면 LOT 열 왼쪽 열 시도
    if not prod_col and lot_col in cols:
        lot_idx = cols.index(lot_col)
        if lot_idx > 0:
            prod_col = cols[lot_idx - 1]

    result = []
    seen = set()
    for _, row in df.iterrows():
        lot = str(row[lot_col]).strip().upper()
        if not lot_pat.match(lot):
            continue
        if lot[0] not in maker_codes:
            continue
        # 10자리 LOT는 끝자리 1자리 제거 (재고장 형식 보정: 9자리가 표준)
        if len(lot) == 10:
            lot = lot[:-1]
        if lot in seen:
            continue
        seen.add(lot)
        product = str(row[prod_col]).strip() if prod_col else ""
        # 품명이 헤더 텍스트이거나 너무 길면 제외
        if product.lower() in ("nan", "none", "") or len(product) > 50:
            product = ""
        result.append({"lot": lot, "product": product})
    return result


def extract_lot_list_from_pdf(pdf_bytes: bytes, filename: str, api_key: str, model: str = "claude-opus-4-8") -> list:
    """PDF/이미지/엑셀에서 LOT번호+품명 목록 추출 (재고 대량 등록용)."""
    import re
    import base64 as _b64

    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    # 엑셀/CSV는 AI 없이 직접 파싱
    if ext in ("xlsx", "xls", "csv"):
        return _extract_lots_from_excel(pdf_bytes, filename)

    client = anthropic.Anthropic(api_key=api_key)

    if ext == "pdf":
        b64 = _b64.standard_b64encode(pdf_bytes).decode("utf-8")
        content = [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": b64}},
            {"type": "text", "text": LOT_EXTRACT_PROMPT},
        ]
    else:
        b64 = encode_image_to_base64(pdf_bytes)
        media_type = detect_media_type(filename)
        content = [
            {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
            {"type": "text", "text": LOT_EXTRACT_PROMPT},
        ]

    response = client.messages.create(
        model=model,
        max_tokens=16000,
        tools=[EXTRACT_LOT_TOOL],
        tool_choice={"type": "tool", "name": "extract_lot_list"},
        messages=[{"role": "user", "content": content}],
    )

    for block in response.content:
        if getattr(block, "type", "") == "tool_use" and block.name == "extract_lot_list":
            items = block.input.get("items", [])
            result = []
            for item in items:
                lot = str(item.get("lot", "")).strip().upper()
                if not re.match(r"^[A-Z]\d{8}$", lot):
                    continue
                product = str(item.get("product", "")).strip()
                result.append({"lot": lot, "product": product})
            return result

    raise ValueError("LOT 추출 응답을 받지 못했습니다.")


def flatten_production_plan(plan_data: dict) -> list:
    """하위 호환: 범용 JSON → flat 리스트"""
    rows = []
    if "items" in plan_data and "lines" not in plan_data:
        for item in plan_data.get("items", []):
            rows.append({
                "라인": item.get("extra_info", ""),
                "위치": "",
                "색상코드": item.get("color_code", ""),
                "제조사": item.get("manufacturer", ""),
                "재고": 0,
                "신규": int(item.get("quantity", 0)),
                "생산량": 0,
            })
        return rows

    for line in plan_data.get("lines", []):
        for item in line.get("items", []):
            rows.append({
                "라인": line.get("line_name", ""),
                "위치": item.get("position", ""),
                "색상코드": item.get("color_code", ""),
                "제조사": item.get("manufacturer", ""),
                "재고": int(item.get("stock", 0)),
                "신규": int(item.get("new_order", item.get("quantity", 0))),
                "생산량": int(item.get("production_qty", 0)),
            })
    return rows
