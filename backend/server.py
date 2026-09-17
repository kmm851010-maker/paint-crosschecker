# =====================================================================
# Project: paint-crosschecker
# Copyright (c) 2026 kmm851010-maker. All rights reserved.
# Unauthorized copying, modification, or distribution is strictly prohibited.
# =====================================================================
"""
페인트 입고 교차검증 시스템 - FastAPI Backend
모바일 앱을 위한 REST API 서버
"""

import base64
import hashlib
import os
import sys

import pandas as pd
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, field_validator

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from modules.vision_ocr import extract_production_plan
from modules.erp_parser import process_erp_file
from modules.matcher import cross_check
from modules.excel_generator import generate_report
from modules.excel_converter import generate_incoming_plan_excel, convert_to_excel, convert_erp_filled_to_excel
from utils.formatter import format_summary

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))
load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))

app = FastAPI(title="페인트 입고 검증 API", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_api_key(api_key: str = "") -> str:
    key = api_key or os.getenv("ANTHROPIC_API_KEY", "")
    if not key:
        raise HTTPException(status_code=400, detail="API 키가 필요합니다.")
    return key


# --- 로그인 ---
STAFF_ACCOUNTS = {
    "admin": {"password": "kgsteel1234", "name": "관리자"},
}

for key, val in os.environ.items():
    if key.startswith("STAFF_"):
        emp_id = key[6:]
        parts = val.split(",", 1)
        STAFF_ACCOUNTS[emp_id] = {"password": parts[0], "name": parts[1] if len(parts) > 1 else emp_id}


class LoginRequest(BaseModel):
    employee_id: str
    password: str


@app.post("/api/login")
async def login(req: LoginRequest):
    # 1) Supabase app_users 테이블 인증 시도
    try:
        from utils.supabase_db import authenticate_app_user
        user = authenticate_app_user(req.employee_id, req.password)
        if user:
            token = hashlib.sha256(f"{req.employee_id}:{req.password}".encode()).hexdigest()[:32]
            return {
                "success": True, "token": token,
                "name": user["name"], "employee_id": req.employee_id,
                "role": user.get("role", "user"),
            }
    except Exception:
        pass
    # 2) 하드코딩 계정 폴백 (admin 등)
    account = STAFF_ACCOUNTS.get(req.employee_id)
    if not account or account["password"] != req.password:
        raise HTTPException(status_code=401, detail="사번 또는 비밀번호가 올바르지 않습니다.")
    token = hashlib.sha256(f"{req.employee_id}:{account['password']}".encode()).hexdigest()[:32]
    return {"success": True, "token": token, "name": account["name"], "employee_id": req.employee_id, "role": "admin"}


@app.get("/health")
async def health_check():
    return {"status": "ok"}


# --- 생산계획서 분석 (입고 예정 리스트) ---

class ParsePlanRequest(BaseModel):
    plan_files: list[str]  # base64 encoded
    plan_filenames: list[str]
    api_key: str = ""


@app.post("/api/parse-plan")
async def parse_plan(req: ParsePlanRequest):
    """생산계획서를 분석하여 입고 예정 품목을 반환합니다."""
    key = get_api_key(req.api_key)

    all_items = []
    all_table_data = []
    for file_b64, filename in zip(req.plan_files, req.plan_filenames):
        plan_bytes = base64.b64decode(file_b64)
        try:
            result = extract_production_plan(plan_bytes, filename, key)
            all_items.extend(result["items"])
            if result.get("table_data"):
                all_table_data.append(result["table_data"])
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"생산계획서 분석 실패: {str(e)}")

    merged_table_data = None
    if all_table_data:
        if len(all_table_data) == 1:
            merged_table_data = all_table_data[0]
        else:
            headers = all_table_data[0]["headers"]
            rows = []
            for td in all_table_data:
                rows.extend(td["rows"])
            merged_table_data = {"headers": headers, "rows": rows}

    return {
        "success": True,
        "items": all_items,
        "count": len(all_items),
        "table_data": merged_table_data,
    }


# --- 교차검증 ---

class CrossCheckRequest(BaseModel):
    plan_files: list[str]
    plan_filenames: list[str]
    erp_file: str
    erp_filename: str
    api_key: str = ""


@app.post("/api/cross-check-multi")
async def run_cross_check_multi(req: CrossCheckRequest):
    """생산계획서 + ERP 교차검증."""
    key = get_api_key(req.api_key)

    # 생산계획서 분석
    all_plan_rows = []
    for file_b64, filename in zip(req.plan_files, req.plan_filenames):
        plan_bytes = base64.b64decode(file_b64)
        try:
            result = extract_production_plan(plan_bytes, filename, key)
            all_plan_rows.extend(result["items"])
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"생산계획서 분석 실패: {str(e)}")

    plan_df = pd.DataFrame(all_plan_rows)

    # ERP 분석
    erp_bytes = base64.b64decode(req.erp_file)
    try:
        erp_df = process_erp_file(erp_bytes, req.erp_filename, key)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ERP 분석 실패: {str(e)}")

    # 교차검증
    result_df = cross_check(plan_df, erp_df)
    summary = format_summary(result_df)

    return {
        "success": True,
        "plan_items": all_plan_rows,
        "erp_items": erp_df.to_dict(orient="records"),
        "results": result_df.to_dict(orient="records"),
        "summary": summary,
    }


# --- 엑셀 다운로드 ---

@app.post("/api/export-excel-multi")
async def export_excel_multi(req: CrossCheckRequest):
    """교차검증 결과 엑셀."""
    key = get_api_key(req.api_key)

    all_plan_rows = []
    for file_b64, filename in zip(req.plan_files, req.plan_filenames):
        plan_bytes = base64.b64decode(file_b64)
        try:
            result = extract_production_plan(plan_bytes, filename, key)
            all_plan_rows.extend(result["items"])
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"분석 실패: {str(e)}")

    plan_df = pd.DataFrame(all_plan_rows)
    erp_bytes = base64.b64decode(req.erp_file)
    erp_df = process_erp_file(erp_bytes, req.erp_filename, key)
    result_df = cross_check(plan_df, erp_df)
    excel_bytes = generate_report(result_df)

    return {
        "success": True,
        "excel_base64": base64.b64encode(excel_bytes).decode("utf-8"),
    }


# --- ERP 입고 반영 결과 ---

class ErpFillRequest(BaseModel):
    table_data: dict  # {"headers": [...], "rows": [[...], ...]}
    erp_results: list[dict]  # [{"색상코드": "...", "입고수량": 5}, ...]


@app.post("/api/erp-fill")
async def erp_fill_endpoint(req: ErpFillRequest):
    """ERP 입고수량을 생산계획서 table_data에 채워 반영합니다."""
    import re as _re
    from utils.helpers import auto_correct_code, is_valid_item_code

    headers = req.table_data.get("headers", [])
    rows_raw = req.table_data.get("rows", [])

    if not headers:
        return {"success": True, "headers": [], "rows": [], "excel_base64": ""}

    # 1. ERP 수량 맵
    erp_qty_map = {
        str(r.get("색상코드", "")).strip(): int(r.get("입고수량", 0) or 0)
        for r in req.erp_results
        if str(r.get("색상코드", "")).strip()
    }

    # 2. 헤더 중복 해소
    unique_h, seen = [], {}
    for h in headers:
        hs = str(h) if h else ""
        if hs in seen:
            seen[hs] += 1
            unique_h.append(f"{hs}_{seen[hs]}")
        else:
            seen[hs] = 0
            unique_h.append(hs)

    padded = [list(r[:len(unique_h)]) + [""] * max(0, len(unique_h) - len(r)) for r in rows_raw]

    # 3. 신규 다음 입고 컬럼 삽입
    exp_h = []
    new_col_flags = []
    for h in unique_h:
        exp_h.append(h)
        new_col_flags.append(False)
        if "신규" in str(h):
            in_name = str(h).replace("신규", "입고")
            base_n, c = in_name, 1
            while in_name in exp_h:
                in_name = f"{base_n}_{c}"; c += 1
            exp_h.append(in_name)
            new_col_flags.append(True)

    def _ts(x):
        if x is None or str(x) in ("nan", "None", "NaN"): return ""
        if isinstance(x, float): return str(int(x)) if x == int(x) else str(x)
        return x if isinstance(x, str) else str(x)

    exp_rows = []
    for rw in padded:
        new_row = []
        orig_idx = 0
        for is_new in new_col_flags:
            if is_new:
                new_row.append("")
            else:
                new_row.append(_ts(rw[orig_idx]) if orig_idx < len(rw) else "")
                orig_idx += 1
        exp_rows.append(new_row)

    # 4. 입고 컬럼에 ERP 수량 채우기 (code는 신규 인덱스 - 3)
    신규_col_indices = [i for i, h in enumerate(exp_h) if "신규" in str(h)]
    for row_idx in range(len(exp_rows)):
        for ni in 신규_col_indices:
            if ni + 1 >= len(exp_h): continue
            inc_col_name = exp_h[ni + 1]
            if "입고" not in str(inc_col_name): continue
            code_col_idx = ni - 3
            if code_col_idx < 0: continue
            raw_code = str(exp_rows[row_idx][code_col_idx]).strip()
            corrected = auto_correct_code(raw_code)
            if is_valid_item_code(corrected) and corrected in erp_qty_map:
                qty = erp_qty_map[corrected]
                exp_rows[row_idx][ni + 1] = str(qty) if qty > 0 else ""

    # 5. 중복 코드 행 병합 (같은 코드 두 번 나오면 신규 합산, 중복 행 입고 클리어)
    for ni in 신규_col_indices:
        if ni + 1 >= len(exp_h): continue
        inc_col_idx = ni + 1
        code_col_idx = ni - 3
        if code_col_idx < 0: continue
        first_seen: dict = {}
        for ridx in range(len(exp_rows)):
            raw = str(exp_rows[ridx][code_col_idx]).strip()
            corr = auto_correct_code(raw)
            if not is_valid_item_code(corr): continue
            if corr not in first_seen:
                first_seen[corr] = ridx
            else:
                fidx = first_seen[corr]
                try: fn = int(str(exp_rows[fidx][ni]).strip() or "0")
                except: fn = 0
                try: dn = int(str(exp_rows[ridx][ni]).strip() or "0")
                except: dn = 0
                if dn > 0:
                    exp_rows[fidx][ni] = str(fn + dn)
                    exp_rows[ridx][ni] = ""
                exp_rows[ridx][inc_col_idx] = ""

    # 6. 재고 위치 조회
    inv_location_map = {}
    try:
        from utils.inventory_supabase import get_sector_inventory
        inv_raw = get_sector_inventory()
        prod_sectors: dict = {}
        for sec, drums in inv_raw.items():
            for d in drums:
                p = str(d.get("product", "")).strip().upper()
                if p:
                    prod_sectors.setdefault(p, {})
                    prod_sectors[p][sec] = prod_sectors[p].get(sec, 0) + 1
        inv_location_map = {
            p: " / ".join(f"{s}({n})" for s, n in sv.items())
            for p, sv in prod_sectors.items()
        }
    except Exception:
        pass

    # 7. 재고 컬럼 우측에 위치 삽입 (역순, code는 재고_idx - 2)
    final_h = list(exp_h)
    final_rows = [list(row) for row in exp_rows]
    if inv_location_map:
        jaego_idxs = [(i, c) for i, c in enumerate(final_h) if _re.match(r'^재고(_\d+)?$', str(c).strip())]
        for rc_idx, rc in reversed(jaego_idxs):
            suffix = str(rc)[2:]
            code_for_loc_idx = rc_idx - 2
            wi_col_name = f"위치{suffix}"
            insert_pos = rc_idx + 1
            loc_vals = []
            for row in final_rows:
                stock = str(row[rc_idx]).strip() if rc_idx < len(row) else ""
                try: stock_n = int(stock) if stock else 0
                except: stock_n = 0
                loc = ""
                if stock_n > 0 and 0 <= code_for_loc_idx < len(row):
                    code = str(row[code_for_loc_idx]).strip().upper()
                    loc = inv_location_map.get(code, "")
                loc_vals.append(loc)
            final_h.insert(insert_pos, wi_col_name)
            for j, row in enumerate(final_rows):
                row.insert(insert_pos, loc_vals[j])

    # 8. 엑셀 생성 (위치 포함, 상태 컬럼 없는 버전)
    excel_bytes = convert_erp_filled_to_excel(final_h, final_rows)

    # 9. 상태 컬럼 삽입 (표시용)
    disp_h = list(final_h)
    disp_rows = [list(row) for row in final_rows]
    orig_h = list(final_h)
    insert_offset = 0
    for fi, fh in enumerate(orig_h):
        if "신규" in str(fh):
            next_i = fi + 1
            if next_i < len(orig_h) and "입고" in str(orig_h[next_i]):
                inc_col = orig_h[next_i]
                insert_at = fi + insert_offset + 2
                st_vals = []
                for row in disp_rows:
                    try:
                        sq_i = disp_h.index(fh)
                        iq_i = disp_h.index(inc_col)
                        sq = int(str(row[sq_i]).strip() or "0") if sq_i < len(row) else 0
                        iq = int(str(row[iq_i]).strip() or "0") if iq_i < len(row) else 0
                    except Exception:
                        sq, iq = 0, 0
                    if sq == 0: st_vals.append("")
                    elif iq == 0: st_vals.append("🟥 미입고")
                    elif iq == sq: st_vals.append("🟩 일치")
                    elif iq > sq: st_vals.append("🟡 초과")
                    else: st_vals.append("🟠 일부")
                disp_h.insert(insert_at, "상태")
                for j, row in enumerate(disp_rows):
                    row.insert(insert_at, st_vals[j])
                insert_offset += 1

    return {
        "success": True,
        "headers": disp_h,
        "rows": disp_rows,
        "excel_base64": base64.b64encode(excel_bytes).decode("utf-8"),
    }


# --- 생산계획서 변환결과 (입고/위치 컬럼 추가) ---

class PlanConversionRequest(BaseModel):
    table_data: dict  # {"headers": [...], "rows": [[...], ...]}
    plan_items: list[dict] = []


@app.post("/api/plan-conversion")
async def plan_conversion_endpoint(req: PlanConversionRequest):
    """생산계획서 변환결과 테이블 생성 (신규→입고 컬럼, 재고→위치 컬럼 추가)."""
    import re as _re

    headers = req.table_data.get("headers", [])
    rows_raw = req.table_data.get("rows", [])

    if not headers:
        return {"success": True, "headers": [], "rows": [], "excel_base64": ""}

    # 1. 헤더 중복 해소
    unique_h, seen = [], {}
    for h in headers:
        hs = str(h) if h else ""
        if hs in seen:
            seen[hs] += 1
            unique_h.append(f"{hs}_{seen[hs]}")
        else:
            seen[hs] = 0
            unique_h.append(hs)

    # 2. 행 패딩
    padded = [list(r[:len(unique_h)]) + [""] * max(0, len(unique_h) - len(r)) for r in rows_raw]

    # 3. 신규 컬럼 다음에 입고 컬럼 삽입
    final_h = []
    new_col_flags = []
    for h in unique_h:
        final_h.append(h)
        new_col_flags.append(False)
        if "신규" in str(h):
            in_name = str(h).replace("신규", "입고")
            base_n, c = in_name, 1
            while in_name in final_h:
                in_name = f"{base_n}_{c}"; c += 1
            final_h.append(in_name)
            new_col_flags.append(True)

    def _ts(x):
        if x is None or str(x) in ("nan", "None", "NaN"):
            return ""
        if isinstance(x, float):
            return str(int(x)) if x == int(x) else str(x)
        return x if isinstance(x, str) else str(x)

    final_rows = []
    for rw in padded:
        new_row = []
        orig_idx = 0
        for is_new in new_col_flags:
            if is_new:
                new_row.append("")
            else:
                new_row.append(_ts(rw[orig_idx]) if orig_idx < len(rw) else "")
                orig_idx += 1
        final_rows.append(new_row)

    # 4. 재고 위치 조회
    inv_location_map = {}
    try:
        from utils.inventory_supabase import get_sector_inventory
        inv_raw = get_sector_inventory()
        prod_sectors: dict = {}
        for sec, drums in inv_raw.items():
            for d in drums:
                p = str(d.get("product", "")).strip().upper()
                if p:
                    prod_sectors.setdefault(p, {})
                    prod_sectors[p][sec] = prod_sectors[p].get(sec, 0) + 1
        inv_location_map = {
            p: " / ".join(f"{s}({n})" for s, n in sv.items())
            for p, sv in prod_sectors.items()
        }
    except Exception:
        pass

    # 5. 색상코드 컬럼 탐지
    code_col_idx = None
    code_kws = ["색상코드", "품목코드", "clrcd", "color", "코드", "품목"]
    for i, ch in enumerate(final_h):
        if any(kw in str(ch).lower() for kw in code_kws):
            code_col_idx = i
            break
    if code_col_idx is None:
        for i, ch in enumerate(final_h):
            vals = [row[i] for row in final_rows if i < len(row)]
            match_cnt = sum(1 for v in vals if _re.match(r"^[A-Za-z][A-Za-z0-9]{5,}$", str(v).strip()))
            if vals and match_cnt >= len(vals) * 0.4:
                code_col_idx = i
                break

    # 6. 재고 컬럼 우측에 위치 컬럼 삽입 (역순으로 처리해 인덱스 유지)
    if inv_location_map:
        jaego_idxs = [(i, c) for i, c in enumerate(final_h) if _re.match(r'^재고(_\d+)?$', str(c).strip())]
        for rc_idx, rc in reversed(jaego_idxs):
            suffix = str(rc)[2:]
            corr_code_col = f"색상코드{suffix}"
            if corr_code_col in final_h:
                corr_code_idx = final_h.index(corr_code_col)
            else:
                corr_code_idx = code_col_idx
            wi_col_name = f"위치{suffix}"
            insert_pos = rc_idx + 1
            loc_vals = []
            for row in final_rows:
                stock = str(row[rc_idx]).strip() if rc_idx < len(row) else ""
                try:
                    stock_n = int(stock) if stock else 0
                except Exception:
                    stock_n = 0
                loc = ""
                if stock_n > 0 and corr_code_idx is not None and corr_code_idx < len(row):
                    code = str(row[corr_code_idx]).strip().upper()
                    loc = inv_location_map.get(code, "")
                loc_vals.append(loc)
            final_h.insert(insert_pos, wi_col_name)
            for j, row in enumerate(final_rows):
                row.insert(insert_pos, loc_vals[j])

    # 7. 엑셀 생성
    excel_bytes = convert_to_excel(final_h, final_rows)

    return {
        "success": True,
        "headers": final_h,
        "rows": final_rows,
        "excel_base64": base64.b64encode(excel_bytes).decode("utf-8"),
    }


# --- 입고 예정 엑셀 ---

class IncomingPlanRequest(BaseModel):
    plan_items: list[dict]


@app.post("/api/generate-incoming-excel")
async def generate_incoming_excel_endpoint(req: IncomingPlanRequest):
    """입고 예정 품목 엑셀."""
    plan_df = pd.DataFrame(req.plan_items)
    try:
        excel_bytes = generate_incoming_plan_excel(plan_df)
        return {
            "success": True,
            "excel_base64": base64.b64encode(excel_bytes).decode("utf-8"),
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"입고 예정 엑셀 생성 실패: {str(e)}")


# --- 재고 관리 ---

class ParseBarcodeRequest(BaseModel):
    raw_text: str



class DrumItem(BaseModel):
    lot: str
    product: str
    maker: str
    scanDisabled: bool = False

    @field_validator("scanDisabled", mode="before")
    @classmethod
    def _parse_scan_disabled(cls, v):
        if isinstance(v, bool):
            return v
        if isinstance(v, str):
            return v == "Y"
        return bool(v)


class InventoryRegisterRequest(BaseModel):
    drums: list[DrumItem]
    sector: str
    remark: str = ""
    skip_existing: bool = False


@app.post("/api/inventory/parse-barcode")
async def parse_barcode_endpoint(req: ParseBarcodeRequest):
    """바코드 텍스트에서 LOT/품명/제조사 파싱"""
    from utils.inventory_supabase import parse_barcode
    result = parse_barcode(req.raw_text)
    if not result:
        raise HTTPException(status_code=400, detail="바코드 형식이 올바르지 않습니다. (최소 16자리 필요)")
    return result


@app.post("/api/inventory/register")
async def inventory_register(req: InventoryRegisterRequest):
    """드럼 목록을 섹터에 등록/이동, 라인입고 시 재고에서 제거"""
    from utils.inventory_supabase import save_drums_to_sector, checkout_drums, CHECKOUT_SECTOR, RETURN_SECTOR
    drums = [d.model_dump() for d in req.drums]
    try:
        if req.sector in (CHECKOUT_SECTOR, RETURN_SECTOR):
            checkout_drums(drums)
            return {"success": True, "count": len(drums), "sector": req.sector, "already_same": [], "moved": len(drums)}
        else:
            result = save_drums_to_sector(drums, req.sector, remark=req.remark, skip_existing=req.skip_existing)
            return {"success": True, "count": len(drums), "sector": req.sector,
                    "already_same": result["already_same"], "moved": result["moved"],
                    "skipped": result.get("skipped", [])}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


class ScanDisabledRequest(BaseModel):
    drums: list[DrumItem]
    disabled: bool


@app.post("/api/inventory/scan-disabled")
async def set_scan_disabled_endpoint(req: ScanDisabledRequest):
    """스캔불가 플래그 설정/해제"""
    from utils.inventory_supabase import set_scan_disabled
    drums = [d.model_dump() for d in req.drums]
    try:
        set_scan_disabled(drums, req.disabled)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "count": len(drums), "disabled": req.disabled}


class ReturnStatusRequest(BaseModel):
    drums: list[DrumItem]
    status: str  # "Y" → 반품대기, "" → 해제


@app.post("/api/inventory/return-status")
async def set_return_status_endpoint(req: ReturnStatusRequest):
    """반품상태 플래그 설정/해제 (섹터 변경 없음)"""
    from utils.inventory_supabase import set_return_status
    drums = [d.model_dump() for d in req.drums]
    try:
        set_return_status(drums, req.status)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "count": len(drums), "status": req.status}


@app.get("/api/version")
async def get_version():
    return {"version": "20260901-b64-fix"}


class UpdateDrumRequest(BaseModel):
    old_lot: str
    new_lot: str
    new_product: str
    new_maker: str
    new_sector: str
    new_remark: str = ""


@app.post("/api/inventory/update-drum")
async def update_drum_endpoint(req: UpdateDrumRequest):
    """드럼 정보 수정 (LOT/품명/제조사/섹터/비고)"""
    from utils.inventory_supabase import update_drum_fields
    try:
        update_drum_fields(req.old_lot, req.new_lot, req.new_product, req.new_maker, req.new_sector, new_remark=req.new_remark)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


@app.get("/api/debug/inventory-sheet")
async def debug_inventory_sheet():
    """재고현황 Supabase 진단용"""
    from utils.inventory_supabase import get_sector_inventory as _gsi
    try:
        sectors = _gsi()
        total = sum(len(v) for v in sectors.values())
        return {"source": "supabase", "sector_count": len(sectors), "drum_count": total}
    except Exception as e:
        return {"source": "supabase", "error": str(e)}


@app.get("/api/inventory/sectors")
async def get_inventory_sectors():
    """섹터별 보관 드럼 현황 조회"""
    from utils.inventory_supabase import get_sector_inventory
    try:
        sectors = get_sector_inventory()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "sectors": sectors}


@app.get("/api/inventory/product-whitelist")
async def get_product_whitelist_endpoint():
    """ERP 입고 등록된 품명 화이트리스트 반환 (모바일 스캔 승인 목록)."""
    from utils.inventory_supabase import get_product_whitelist
    try:
        products = get_product_whitelist()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "products": products}


@app.get("/api/inventory/history")
async def get_inventory_history_endpoint(from_dt: str = "", to_dt: str = ""):
    """재고 이력 조회 (from_dt/to_dt: 'YYYY-MM-DD HH:MM', 기본값 오늘 KST 00:00~23:59)"""
    from utils.inventory_supabase import get_inventory_history
    import datetime as _dt
    _today = (_dt.datetime.utcnow() + _dt.timedelta(hours=9)).strftime("%Y-%m-%d")
    if not from_dt:
        from_dt = f"{_today} 00:00"
    if not to_dt:
        to_dt = f"{_today} 23:59"
    try:
        history = get_inventory_history(from_dt, to_dt)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "from_dt": from_dt, "to_dt": to_dt, "history": history}


class ParseReturnListRequest(BaseModel):
    file_data: str  # base64
    filename: str
    api_key: str = ""


class ParsePDFLotsRequest(BaseModel):
    file_data: str  # base64
    filename: str
    api_key: str = ""


@app.post("/api/inventory/parse-pdf-lots")
async def parse_pdf_lots_endpoint(req: ParsePDFLotsRequest):
    """PDF/이미지에서 LOT번호+품명 추출 (재고 대량 등록용)"""
    from modules.vision_ocr import extract_lot_list_from_pdf
    key = get_api_key(req.api_key)
    file_bytes = base64.b64decode(req.file_data)
    try:
        items = extract_lot_list_from_pdf(file_bytes, req.filename, key)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "count": len(items), "items": items}


@app.post("/api/inventory/parse-return-list")
async def parse_return_list_endpoint(req: ParseReturnListRequest):
    """반품 리스트 이미지/엑셀에서 품명·LOT-NO·반품유형 추출"""
    from utils.return_list_parser import parse_return_list_excel, parse_return_list_image
    key = get_api_key(req.api_key)
    file_bytes = base64.b64decode(req.file_data)
    ext = req.filename.lower().rsplit(".", 1)[-1]
    try:
        if ext in ("xlsx", "xls", "csv"):
            items = parse_return_list_excel(file_bytes, ext)
        else:
            items = parse_return_list_image(file_bytes, req.filename, key)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True, "count": len(items), "items": items}


# ═══════════════════════════════════════════════════════════════════
# 직원 관리
# ═══════════════════════════════════════════════════════════════════

class RegisterEmployeeRequest(BaseModel):
    department: str
    name: str
    employee_id: str
    password: str = ""


class ResetPasswordRequest(BaseModel):
    new_password: str


@app.get("/api/employees")
async def get_employees(department: str = ""):
    from utils.supabase_db import list_app_users
    return {"success": True, "employees": list_app_users(department or None)}


@app.post("/api/employees")
async def create_employee(req: RegisterEmployeeRequest):
    from utils.supabase_db import register_app_user
    init_pw = req.password.strip() or req.employee_id
    try:
        register_app_user(req.department, req.name, req.employee_id, init_pw)
    except ValueError as e:
        raise HTTPException(status_code=409, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


@app.delete("/api/employees/{employee_id}")
async def remove_employee(employee_id: str):
    from utils.supabase_db import delete_app_user
    try:
        delete_app_user(employee_id)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


@app.post("/api/employees/{employee_id}/reset-password")
async def reset_employee_password(employee_id: str, req: ResetPasswordRequest):
    from utils.supabase_db import reset_app_user_password
    if not req.new_password.strip():
        raise HTTPException(status_code=400, detail="비밀번호를 입력하세요.")
    try:
        reset_app_user_password(employee_id, req.new_password.strip())
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


class ChangePasswordRequest(BaseModel):
    employee_id: str
    current_password: str
    new_password: str

@app.post("/api/auth/change-password")
async def change_password(req: ChangePasswordRequest):
    from utils.supabase_db import authenticate_app_user, reset_app_user_password
    employee_id = req.employee_id.strip()
    current_pw  = req.current_password
    new_pw      = req.new_password.strip()
    if not employee_id or not current_pw or not new_pw:
        raise HTTPException(status_code=400, detail="필수 항목이 누락되었습니다.")
    if len(new_pw) < 4:
        raise HTTPException(status_code=400, detail="비밀번호는 4자 이상이어야 합니다.")
    user = authenticate_app_user(employee_id, current_pw)
    if not user:
        raise HTTPException(status_code=401, detail="현재 비밀번호가 올바르지 않습니다.")
    reset_app_user_password(employee_id, new_pw)
    return {"success": True}


@app.get("/api/members")
async def get_members(department: str = "칼라반지게차"):
    from utils.supabase_db import get_members_dict
    return {"success": True, "members": get_members_dict(department)}


# ═══════════════════════════════════════════════════════════════════
# 작업일지
# ═══════════════════════════════════════════════════════════════════

class SaveWorklogRequest(BaseModel):
    date: str           # YYYY-MM-DD
    shift_data: dict
    work_items: list
    safety_items: list = []
    note: str = ""
    leave_list: list = []


@app.get("/api/worklog")
async def get_worklog(date: str):
    """date=YYYY-MM-DD. 저장된 작업일지 + 자동 shift 계산 반환."""
    import datetime as _dt
    from utils.supabase_db import (
        load_work_items, load_daily_detail, load_leaves,
        get_monthly_totals, get_members_dict, get_shift_info, apply_leaves,
    )
    try:
        d = _dt.date.fromisoformat(date)
    except ValueError:
        raise HTTPException(status_code=400, detail="날짜 형식이 올바르지 않습니다. (YYYY-MM-DD)")

    members = get_members_dict()
    shift_auto = get_shift_info(d, members)
    leave_list = load_leaves()
    shift_auto = apply_leaves(shift_auto, d, leave_list)

    detail = load_daily_detail(date)
    work_items = load_work_items(date)
    monthly_totals = get_monthly_totals(date)

    return {
        "success": True,
        "date": date,
        "members": members,
        "shift_auto": shift_auto,
        "saved_shift": (detail.get("shift") if detail else None),
        "saved_safety": (detail.get("safety") if detail else []),
        "saved_note": (detail.get("note") if detail else ""),
        "work_items": work_items,
        "monthly_totals": monthly_totals,
        "leave_list": leave_list,
    }


@app.post("/api/worklog")
async def save_worklog(req: SaveWorklogRequest):
    from utils.supabase_db import save_work_items, save_daily_detail
    try:
        save_work_items(req.date, req.work_items)
        save_daily_detail(req.date, req.shift_data, req.safety_items, req.note)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


_WORKLOG_ITEM_NAMES = [
    "페인트 하차 수량", "페인트 공급 수량", "재고 페인트 창고 입고",
    "AGV 입/출고 작업 수량",
    "신나 하차 수량", "신나 공급 수량", "크롬 공급 수량",
    "공드럼 운반 수량", "페보루 운반 수량", "페신너 운반 및 상차",
    "반품 , 불량 페인트 수량", "코터롤 운반 횟수", "필름 하차, 장소 이동 횟수",
]


def _build_worklog_sheet(ws, selected_date, shift_data: dict, work_items: list, safety_items: list, note_text: str):
    """워크시트에 일일 작업일지를 채운다 (10컬럼 A-J). Streamlit _fill_work_log_sheet 동일 포맷."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    font_title = Font(name="맑은 고딕", size=18, bold=True, underline="single")
    font_date  = Font(name="맑은 고딕", size=11, bold=True)
    font_sec   = Font(name="맑은 고딕", size=11, bold=True)
    font_hdr   = Font(name="맑은 고딕", size=9, bold=True)
    font_body  = Font(name="맑은 고딕", size=9)
    font_bold  = Font(name="맑은 고딕", size=9, bold=True)
    fill_gray  = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    _t = Side(style="thin", color="000000")
    thin = Border(left=_t, right=_t, top=_t, bottom=_t)
    align_c  = Alignment(horizontal="center", vertical="center")
    align_cw = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, w in zip("ABCDEFGHIJ", [9, 5, 13, 9, 9, 9, 9, 9, 9, 10]):
        ws.column_dimensions[col].width = w

    # 제목
    ws.row_dimensions[2].height = 32
    ws.merge_cells("A2:J2")
    import os
    team = os.environ.get("TEAM_NAME", "")
    ws["A2"] = f"{team} 일일 업무 보고" if team else "일일 업무 보고"
    ws["A2"].font, ws["A2"].alignment = font_title, align_c

    # 날짜
    ws.merge_cells("H4:J4")
    ws["H4"] = selected_date.strftime("%Y년 %m월 %d일")
    ws["H4"].font = font_date
    ws["H4"].alignment = Alignment(horizontal="right", vertical="center")

    # 1. 인원 현황
    ws["A5"] = "1. 인원 현황"
    ws["A5"].font = font_sec
    for pos, txt in [("A6","구분"),("B6","조"),("C6","근무시간"),("D6","근무자"),("E6","휴무자"),("F6","연장근무")]:
        ws[pos] = txt
        ws[pos].font, ws[pos].fill, ws[pos].alignment, ws[pos].border = font_hdr, fill_gray, align_c, thin
    ws.merge_cells("G6:J6")
    ws["G6"] = "비고"
    ws["G6"].font, ws["G6"].fill, ws["G6"].alignment, ws["G6"].border = font_hdr, fill_gray, align_c, thin
    for c in "HIJ":
        ws[f"{c}6"].fill, ws[f"{c}6"].border = fill_gray, thin

    if shift_data.get("is_allleave"):
        fill_purple = PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")
        font_purple = Font(name="맑은 고딕", size=11, bold=True, color="5B21B6")
        ws.merge_cells("A7:J7")
        ws["A7"] = f"전원 {shift_data.get('leave_type', '명휴')} — 해당일 전원 휴무"
        ws["A7"].font = font_purple
        ws["A7"].fill = fill_purple
        ws["A7"].alignment = align_c
        ws["A7"].border = thin
        rows_1 = []  # 전원 명휴: 인원 행 없음
    elif shift_data.get("is_2person"):
        rows_1 = [
            ("주간", shift_data.get("주간_조", shift_data.get("1근_조","")), "06:30 – 18:30",
             shift_data.get("주간_근무자", shift_data.get("1근_근무자","")), "", shift_data.get("1근_연장",""), shift_data.get("1근_비고","")),
            ("야간", shift_data.get("야간_조", shift_data.get("2근_조","")), "18:30 – 06:30",
             shift_data.get("야간_근무자", shift_data.get("2근_근무자","")), "", shift_data.get("2근_연장",""), shift_data.get("2근_비고","")),
            ("휴무", "", "", "", shift_data.get("3근_근무자",""), "", shift_data.get("3근_비고","")),
            ("휴무", shift_data.get("휴무_조",""), "", "", shift_data.get("휴무_근무자",""), "", shift_data.get("휴무_구분",""))
        ]
    else:
        # 부분 휴가 처리: leave_person에 있는 근무자는 휴무자 칸으로 이동, 비고에 휴가 유형 표시
        absent_set = {n.strip() for n in shift_data.get("leave_person", "").split(",") if n.strip()}
        ltype = shift_data.get("leave_type", "")

        def _row(label, team, time_, worker, ext, note):
            if worker in absent_set:
                return (label, team, time_, "", worker, ext, ltype or note)
            return (label, team, time_, worker, "", ext, note)

        rows_1 = [
            _row("1근", shift_data.get("1근_조",""), "06:30 – 14:30", shift_data.get("1근_근무자",""), shift_data.get("1근_연장",""), shift_data.get("1근_비고","")),
            _row("2근", shift_data.get("2근_조",""), "14:30 – 22:30", shift_data.get("2근_근무자",""), shift_data.get("2근_연장",""), shift_data.get("2근_비고","")),
            _row("3근", shift_data.get("3근_조",""), "22:30 – 06:30", shift_data.get("3근_근무자",""), shift_data.get("3근_연장",""), shift_data.get("3근_비고","")),
            ("휴무", shift_data.get("휴무_조",""), "", "", shift_data.get("휴무_근무자",""), "", shift_data.get("휴무_구분",""))
        ]
    for ri, r in enumerate(rows_1, start=7):
        ws.row_dimensions[ri].height = 28
        for ci, val in zip("ABCDEF", r[:6]):
            if ci == "F" and val == 0:
                val = ""
            ws[f"{ci}{ri}"] = val
            ws[f"{ci}{ri}"].font  = font_body
            ws[f"{ci}{ri}"].border = thin
            ws[f"{ci}{ri}"].alignment = align_cw if ci == "C" else align_c
        ws.merge_cells(f"G{ri}:J{ri}")
        ws[f"G{ri}"] = r[6]
        for c in "GHIJ":
            ws[f"{c}{ri}"].font, ws[f"{c}{ri}"].border = font_body, thin
        ws[f"G{ri}"].alignment = align_c

    # 2. 업무 현황
    ws["A12"] = "2. 업무 현황"
    ws["A12"].font = font_sec
    ws.merge_cells("A13:C13")
    ws["A13"] = "작업 내용"
    ws["A13"].font, ws["A13"].fill, ws["A13"].alignment, ws["A13"].border = font_hdr, fill_gray, align_c, thin
    for c in "BC":
        ws[f"{c}13"].fill, ws[f"{c}13"].border = fill_gray, thin
    for col_l, hdr in zip("DEFGHIJ", ["1근","2근","3근","주간","야간","합계","월합계"]):
        ws[f"{col_l}13"] = hdr
        ws[f"{col_l}13"].font, ws[f"{col_l}13"].fill, ws[f"{col_l}13"].alignment, ws[f"{col_l}13"].border = font_hdr, fill_gray, align_c, thin

    # work_items: 이름 → dict 변환
    item_map = {it["name"]: it for it in (work_items or [])}
    for wi, name in enumerate(_WORKLOG_ITEM_NAMES, start=14):
        it = item_map.get(name, {})
        ws.row_dimensions[wi].height = 20
        ws.merge_cells(f"A{wi}:C{wi}")
        ws[f"A{wi}"] = name
        ws[f"A{wi}"].font, ws[f"A{wi}"].alignment, ws[f"A{wi}"].border = font_bold, align_c, thin
        for c in "BC":
            ws[f"{c}{wi}"].border = thin
        ws[f"D{wi}"] = it.get("s1") or ""
        ws[f"E{wi}"] = it.get("s2") or ""
        ws[f"F{wi}"] = it.get("s3") or ""
        ws[f"G{wi}"] = it.get("day") or ""
        ws[f"H{wi}"] = it.get("night") or ""
        for c in "DEFGH":
            ws[f"{c}{wi}"].font, ws[f"{c}{wi}"].alignment, ws[f"{c}{wi}"].border = font_body, align_c, thin
        ws[f"I{wi}"] = f"=SUM(D{wi}:H{wi})"
        ws[f"I{wi}"].font, ws[f"I{wi}"].alignment, ws[f"I{wi}"].border = font_bold, align_c, thin
        ws[f"J{wi}"] = it.get("month_total") or 0
        ws[f"J{wi}"].font, ws[f"J{wi}"].alignment, ws[f"J{wi}"].border = font_body, align_c, thin

    # 3. 안전 관리 사항
    ws["A27"] = "3. 안전 관리 사항"
    ws["A27"].font = font_sec
    ws.merge_cells("A28:E28")
    for c in "ABCDE":
        ws[f"{c}28"].fill, ws[f"{c}28"].border = fill_gray, thin
    for pos, txt in [("F28","1근"),("G28","2근"),("H28","3근"),("I28","주간"),("J28","야간")]:
        ws[pos] = txt
        ws[pos].font, ws[pos].fill, ws[pos].alignment, ws[pos].border = font_hdr, fill_gray, align_c, thin

    align_safety = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for si, s_row in enumerate(safety_items or [], start=29):
        ws.row_dimensions[si].height = 34
        ws.merge_cells(f"A{si}:E{si}")
        ws[f"A{si}"] = s_row.get("text", "")
        ws[f"A{si}"].font, ws[f"A{si}"].alignment, ws[f"A{si}"].border = font_bold, align_safety, thin
        for c in "BCDE":
            ws[f"{c}{si}"].border = thin
        for pc, key in [("F","s1"),("G","s2"),("H","s3"),("I","day"),("J","night")]:
            cell = ws[f"{pc}{si}"]
            cell.value = "☑" if s_row.get(key) else "□"
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment, cell.border = align_c, thin

    # 4. 특이 사항
    ws["A36"] = "4. 특이 사항"
    ws["A36"].font = font_sec
    ws.merge_cells("A37:J41")
    ws["A37"] = note_text or ""
    ws["A37"].font = font_body
    ws["A37"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for r in range(37, 42):
        for c_idx in range(1, 11):
            cl = get_column_letter(c_idx)
            ws[f"{cl}{r}"].border = Border(
                top=_t if r == 37 else Side(),
                bottom=_t if r == 41 else Side(),
                left=_t if c_idx == 1 else Side(),
                right=_t if c_idx == 10 else Side()
            )


@app.get("/api/worklog/export")
async def export_worklog_excel(year: int, month: int, day: int = 0):
    """월간 작업일지 Excel 생성. base64 인코딩으로 반환. day 지정 시 해당 시트를 active로 설정."""
    import io
    import base64
    import calendar as _cal
    import datetime as _dt
    import openpyxl
    from utils.supabase_db import (
        load_daily_detail_month, load_work_items_month,
        get_members_dict, get_shift_info, apply_leaves, load_leaves,
    )

    detail_by_date = load_daily_detail_month(year, month)
    work_by_date = load_work_items_month(year, month)
    members = get_members_dict()
    leave_list = load_leaves()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    last_day = _cal.monthrange(year, month)[1]
    added = 0
    for d in range(1, last_day + 1):
        date_obj = _dt.date(year, month, d)
        date_str = date_obj.strftime("%Y-%m-%d")
        detail = detail_by_date.get(date_str) or {}
        shift = detail.get("shift") or {}
        if not shift:
            shift = get_shift_info(date_obj, members)
        # saved_shift 여부와 무관하게 항상 최신 휴가 정보 적용 (leave_person 등 동기화)
        shift = apply_leaves(shift, date_obj, leave_list)
        ws_new = wb.create_sheet(title=f"{month}월{d}일")
        _build_worklog_sheet(
            ws_new, date_obj,
            shift,
            work_by_date.get(date_str, []),
            detail.get("safety") or [],
            detail.get("note") or "",
        )
        added += 1

    if added == 0:
        ws_empty = wb.create_sheet(title="데이터없음")
        ws_empty["A1"] = "저장된 작업일지가 없습니다."

    # 요청된 날짜 시트를 active로 설정
    target_title = f"{month}월{day}일" if day else None
    if target_title and target_title in wb.sheetnames:
        wb.active = wb[target_title]
    elif wb.sheetnames:
        wb.active = wb[wb.sheetnames[-1]]  # 없으면 마지막 시트

    buf = io.BytesIO()
    wb.save(buf)
    encoded = base64.b64encode(buf.getvalue()).decode()
    import os
    team = os.environ.get("TEAM_NAME", "")
    filename = f"{year}{month:02d}{team}_작업일지.xlsx"
    return {"success": True, "data": encoded, "count": added, "filename": filename}


class SendEmailRequest(BaseModel):
    year: int
    month: int
    to: str       # 쉼표 구분 이메일
    subject: str
    body: str
    extra_files: list = []  # [{name: str, data: str(base64)}]


@app.post("/api/worklog/send-email")
async def send_worklog_email(req: SendEmailRequest):
    """월간 작업일지 Excel을 첨부해 Gmail API로 메일 전송."""
    import io, base64 as _b64, calendar as _cal, datetime as _dt, os
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.base import MIMEBase
    from email import encoders
    from email.header import Header
    import google.oauth2.credentials as _goauth
    import googleapiclient.discovery as _gdisco
    import openpyxl
    from utils.supabase_db import (
        load_daily_detail_month, load_work_items_month,
        get_members_dict, get_shift_info, apply_leaves, load_leaves,
    )

    # ── Gmail OAuth 설정 읽기 (워크로그 스프레드시트 gmail_config 시트) ──
    sheet_err = ""
    cfg: dict = {}
    ws_id = os.environ.get("WORKLOG_SPREADSHEET_ID", "")
    if ws_id:
        try:
            from utils.inventory_sheets import _get_client
            client = _get_client()
            sp = client.open_by_key(ws_id)
            ws_cfg = sp.worksheet("gmail_config")
            cfg = {r[0]: r[1] for r in ws_cfg.get_all_values() if len(r) >= 2}
        except Exception as _e:
            sheet_err = str(_e)
    # 환경변수 fallback
    if not cfg.get("refresh_token"):
        cfg = {
            "refresh_token": os.environ.get("GMAIL_REFRESH_TOKEN", ""),
            "client_id": os.environ.get("GMAIL_CLIENT_ID", ""),
            "client_secret": os.environ.get("GMAIL_CLIENT_SECRET", ""),
        }

    missing = [k for k in ("refresh_token", "client_id", "client_secret") if not cfg.get(k)]
    if missing:
        parts = [f"누락된 설정: {missing}"]
        if sheet_err:
            parts.append(f"시트 오류: {sheet_err}")
        env_check = {
            "WORKLOG_SPREADSHEET_ID": bool(ws_id),
            "GMAIL_REFRESH_TOKEN": bool(os.environ.get("GMAIL_REFRESH_TOKEN")),
            "GMAIL_CLIENT_ID": bool(os.environ.get("GMAIL_CLIENT_ID")),
            "GMAIL_CLIENT_SECRET": bool(os.environ.get("GMAIL_CLIENT_SECRET")),
        }
        parts.append(f"환경변수 확인: {env_check}")
        raise HTTPException(status_code=500, detail=" | ".join(parts))

    recipients = [r.strip() for r in req.to.split(",") if r.strip()]
    if not recipients:
        raise HTTPException(status_code=400, detail="받는 사람 이메일을 입력하세요.")

    # ── Excel 생성 ──
    year, month = req.year, req.month
    detail_by_date = load_daily_detail_month(year, month)
    work_by_date   = load_work_items_month(year, month)
    members        = get_members_dict()
    leave_list     = load_leaves()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    last_day = _cal.monthrange(year, month)[1]
    for d in range(1, last_day + 1):
        date_obj = _dt.date(year, month, d)
        date_str = date_obj.strftime("%Y-%m-%d")
        detail   = detail_by_date.get(date_str) or {}
        shift    = detail.get("shift") or {}
        if not shift:
            shift = get_shift_info(date_obj, members)
        shift = apply_leaves(shift, date_obj, leave_list)
        ws_new = wb.create_sheet(title=f"{month}월{d}일")
        _build_worklog_sheet(ws_new, date_obj, shift, work_by_date.get(date_str, []), detail.get("safety") or [], detail.get("note") or "")

    if not wb.sheetnames:
        wb.create_sheet(title="데이터없음")["A1"] = "저장된 작업일지가 없습니다."
    if wb.sheetnames:
        wb.active = wb[wb.sheetnames[-1]]

    buf = io.BytesIO()
    wb.save(buf)
    excel_bytes = buf.getvalue()
    team     = os.environ.get("TEAM_NAME", "")
    filename = f"{year}{month:02d}{team}_작업일지.xlsx"

    # ── 메일 구성 ──
    msg = MIMEMultipart("mixed")
    msg["To"]      = ", ".join(recipients)
    msg["Subject"] = Header(req.subject, "utf-8").encode()
    msg.attach(MIMEText(req.body, "plain", "utf-8"))

    part = MIMEBase("application", "octet-stream")
    part.set_payload(excel_bytes)
    encoders.encode_base64(part)
    part.add_header("Content-Disposition", "attachment", filename=("utf-8", "", filename))
    msg.attach(part)

    # ── 추가 첨부파일 ──
    for ef in req.extra_files:
        ef_name = ef.get("name", "file") if isinstance(ef, dict) else getattr(ef, "name", "file")
        ef_data = ef.get("data", "") if isinstance(ef, dict) else getattr(ef, "data", "")
        ef_bytes = _b64.b64decode(ef_data)
        part2 = MIMEBase("application", "octet-stream")
        part2.set_payload(ef_bytes)
        encoders.encode_base64(part2)
        part2.add_header("Content-Disposition", "attachment", filename=("utf-8", "", ef_name))
        msg.attach(part2)

    # ── Gmail API 전송 ──
    creds = _goauth.Credentials(
        token=None,
        refresh_token=cfg["refresh_token"],
        token_uri="https://oauth2.googleapis.com/token",
        client_id=cfg["client_id"],
        client_secret=cfg["client_secret"],
    )
    svc = _gdisco.build("gmail", "v1", credentials=creds)
    raw = _b64.urlsafe_b64encode(msg.as_bytes()).decode()
    svc.users().messages().send(userId="me", body={"raw": raw}).execute()

    return {"success": True, "message": f"메일 전송 완료 → {req.to}"}


# ═══════════════════════════════════════════════════════════════════
# 휴가/대근
# ═══════════════════════════════════════════════════════════════════

class SaveLeavesRequest(BaseModel):
    leave_list: list


@app.get("/api/leaves")
async def get_leaves():
    from utils.supabase_db import load_leaves
    return {"success": True, "leave_list": load_leaves()}


@app.post("/api/leaves")
async def save_leaves_endpoint(req: SaveLeavesRequest):
    from utils.supabase_db import save_leaves
    try:
        save_leaves(req.leave_list)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


# ═══════════════════════════════════════════════════════════════════
# 근무메모 (schedule_notes)
# ═══════════════════════════════════════════════════════════════════

class SaveScheduleNoteRequest(BaseModel):
    name: str
    date: str
    note: str


@app.get("/api/schedule-notes")
async def get_schedule_notes(name: str, year: int, month: int):
    from utils.supabase_db import load_schedule_notes_month
    return {"success": True, "notes": load_schedule_notes_month(name, year, month)}


@app.post("/api/schedule-notes")
async def save_schedule_note_endpoint(req: SaveScheduleNoteRequest):
    from utils.supabase_db import save_schedule_note
    try:
        save_schedule_note(req.name, req.date, req.note)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


# ═══════════════════════════════════════════════════════════════════
# 일일 재고기록
# ═══════════════════════════════════════════════════════════════════

class UpsertRemarkRequest(BaseModel):
    date: str
    shift: str
    product: str
    remark: str


@app.get("/api/daily-inventory")
async def get_daily_inventory(date: str):
    """date=YYYY-MM-DD. 근별 입고 품목 + 비고 반환."""
    import datetime as _dt
    from utils.supabase_db import (
        load_daily_detail, get_inventory_registered_in_range, get_daily_inventory_remarks,
    )
    try:
        d = _dt.date.fromisoformat(date)
    except ValueError:
        raise HTTPException(status_code=400, detail="날짜 형식이 올바르지 않습니다.")

    next_date = (d + _dt.timedelta(days=1)).strftime("%Y-%m-%d")
    detail = load_daily_detail(date)
    shift_data = (detail.get("shift") or {}) if detail else {}

    if not shift_data:
        return {"success": True, "date": date, "shift_data": {}, "shift_groups": []}

    is_2p = shift_data.get("is_2person", False)
    if is_2p:
        shifts = [
            ("주간", f"{date} 06:30:00", f"{date} 18:30:00", shift_data.get("주간_근무자", "")),
            ("야간", f"{date} 18:30:00", f"{next_date} 06:30:00", shift_data.get("야간_근무자", "")),
        ]
    else:
        shifts = [
            ("1근", f"{date} 06:30:00", f"{date} 14:30:00", shift_data.get("1근_근무자", "")),
            ("2근", f"{date} 14:30:00", f"{date} 22:30:00", shift_data.get("2근_근무자", "")),
            ("3근", f"{date} 22:30:00", f"{next_date} 06:30:00", shift_data.get("3근_근무자", "")),
        ]

    remarks_raw = get_daily_inventory_remarks(date)
    remarks_map = {(r["shift"], r["product"]): r["remark"] for r in remarks_raw}

    shift_groups = []
    for sname, sstart, send, sworker in shifts:
        items = get_inventory_registered_in_range(sstart, send)
        pmap: dict = {}
        for it in items:
            if (it.get("remark") or "") == "신규":
                continue
            prod = (it.get("product") or "").strip() or "미상"
            pmap.setdefault(prod, []).append((it.get("lot") or "").strip())
        rows = [
            {"product": p, "qty": len(ls), "lots": sorted(ls),
             "remark": remarks_map.get((sname, p), "")}
            for p, ls in pmap.items()
        ]
        shift_groups.append({"shift": sname, "worker": sworker, "rows": rows})

    return {"success": True, "date": date, "shift_data": shift_data, "shift_groups": shift_groups}


@app.post("/api/daily-inventory/remark")
async def upsert_remark(req: UpsertRemarkRequest):
    from utils.supabase_db import upsert_daily_inventory_remark
    try:
        upsert_daily_inventory_remark(req.date, req.shift, req.product, req.remark)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return {"success": True}


class DailyInventoryExportRequest(BaseModel):
    date: str
    shift_groups: list[dict]


@app.post("/api/daily-inventory/export")
async def export_daily_inventory(req: DailyInventoryExportRequest):
    """일일 재고기록 서식 적용 엑셀 생성."""
    import base64
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    PURPLE = PatternFill(start_color="4B2D8E", end_color="4B2D8E", fill_type="solid")
    GRAY   = PatternFill(start_color="DFE4EA", end_color="DFE4EA", fill_type="solid")
    LIGHT  = PatternFill(start_color="F5F3FF", end_color="F5F3FF", fill_type="solid")
    W_FONT = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    D_FONT = Font(name="맑은 고딕", size=10)
    B_FONT = Font(name="맑은 고딕", size=10, bold=True)
    thin   = Side(style="thin", color="AAAAAA")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    C      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    L      = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    wb = Workbook()
    wb.remove(wb.active)

    for group in req.shift_groups:
        shift   = group.get("shift", "")
        worker  = group.get("worker", "")
        rows    = group.get("rows", [])
        if not rows:
            continue

        ws = wb.create_sheet(title=shift)
        ws.freeze_panes = "A3"

        # 1행: 날짜·근무 정보
        title_text = f"{req.date}  {shift}" + (f"  ({worker})" if worker else "")
        ws.merge_cells("A1:E1")
        t = ws["A1"]
        t.value        = title_text
        t.font         = Font(name="맑은 고딕", size=12, bold=True, color="FFFFFF")
        t.fill         = PURPLE
        t.alignment    = C
        t.border       = BORDER
        ws.row_dimensions[1].height = 22

        # 2행: 헤더
        for ci, h in enumerate(["No.", "품명", "수량", "LOT번호", "비고"], 1):
            c = ws.cell(row=2, column=ci, value=h)
            c.font = W_FONT; c.fill = PURPLE; c.alignment = C; c.border = BORDER
        ws.row_dimensions[2].height = 18

        # 데이터
        r = 3
        total_qty = 0
        no = 1
        for row in rows:
            product = row.get("product", "")
            qty     = row.get("qty", 0)
            lots    = sorted(row.get("lots", [])) or [""]
            remark  = row.get("remark", "")
            total_qty += qty
            fill = LIGHT if no % 2 == 0 else None

            for li, lot in enumerate(lots):
                for ci in range(1, 6):
                    c = ws.cell(row=r + li, column=ci)
                    c.border = BORDER
                    c.font   = D_FONT
                    if fill:
                        c.fill = fill

                ws.cell(row=r + li, column=4, value=lot).alignment = C

            # 첫 행에만 No./품명/수량/비고
            ws.cell(row=r, column=1, value=no).alignment = C
            ws.cell(row=r, column=2, value=product).alignment = L
            ws.cell(row=r, column=3, value=qty).alignment = C
            ws.cell(row=r, column=5, value=remark).alignment = L

            # 여러 LOT이면 병합
            if len(lots) > 1:
                for col in [1, 2, 3, 5]:
                    ws.merge_cells(
                        start_row=r, start_column=col,
                        end_row=r + len(lots) - 1, end_column=col
                    )
                    ws.cell(row=r, column=col).alignment = Alignment(
                        horizontal=("center" if col in [1, 3] else "left"),
                        vertical="center", wrap_text=True
                    )

            r += len(lots)
            no += 1

        # 합계 행
        for ci in range(1, 6):
            c = ws.cell(row=r, column=ci)
            c.fill = GRAY; c.font = B_FONT; c.border = BORDER; c.alignment = C
        ws.cell(row=r, column=2, value="합  계").alignment = C
        ws.cell(row=r, column=3, value=total_qty)

        # 열 너비
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 7
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 20

    buf = BytesIO()
    wb.save(buf)
    b64 = base64.b64encode(buf.getvalue()).decode()
    return {"success": True, "excel_base64": b64}


# ═══════════════════════════════════════════════════════════════════
# 근태관리 — 공휴일 조회
# ═══════════════════════════════════════════════════════════════════

_KG_COMPANY_DAYS = {(9, 1): "창립기념일"}  # 회사 기념일 (월, 일)


@app.get("/api/attendance/holidays")
async def get_holidays(year: int):
    """한국 법정공휴일 + KG 회사 기념일 반환 (연도 기준)"""
    try:
        import holidays as _hol
        kr = _hol.SouthKorea(years=[year - 1, year, year + 1])
        result = {d.strftime("%Y-%m-%d"): name for d, name in kr.items() if d.year == year}
    except Exception:
        result = {}
    for (m, d), name in _KG_COMPANY_DAYS.items():
        result[f"{year}-{m:02d}-{d:02d}"] = name
    return {"holidays": result}


# ═══════════════════════════════════════════════════════════════════
# 근태관리 — 급여시간표 / 교대주기별 연장
# ═══════════════════════════════════════════════════════════════════

@app.get("/api/attendance/month-stats")
async def get_attendance_month_stats(year: int, month: int, name: str):
    """
    특정 근무자의 월별 급여시간표 및 교대주기별 연장 시간 반환.
    year: 2026, month: 1-12, name: 근무자 이름
    """
    import datetime as _dt
    import calendar as _cal
    from utils.supabase_db import (
        get_members_dict, load_daily_detail_month, load_leaves,
    )

    CYCLE_20 = [
        ('B','C','D','A'),('B','C','A','D'),('B','C','A','D'),
        ('B','D','A','C'),('B','D','A','C'),('C','D','A','B'),
        ('C','D','B','A'),('C','D','B','A'),('C','A','B','D'),
        ('C','A','B','D'),('D','A','B','C'),('D','A','C','B'),
        ('D','A','C','B'),('D','B','C','A'),('D','B','C','A'),
        ('A','B','C','D'),('A','B','D','C'),('A','B','D','C'),
        ('A','C','D','B'),('A','C','D','B'),
    ]
    BASE_DATE = _dt.date(2026, 3, 1)
    NIGHT_HOURS = {"1근": 0, "2근": 0.5, "3근": 7.5, "주간": 0, "야간": 7.5}
    SCOLS = ["정상근로","유휴근로","휴일근로","연장근로","휴일연장","야간근로",
             "휴일비근로","휴가비근로","스틸아카데미","항군교육",
             "사내교육(1)","사내교육(1.5)","사외교육(1)","사외교육(1.5)","공가"]

    try:
        members = get_members_dict() or {}
        daily_details = load_daily_detail_month(year, month)
        base_leaves = load_leaves()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

    # 한국 공휴일
    try:
        import holidays as _hol
        kr_holidays = _hol.SouthKorea(years=[year, year-1, year+1])
    except Exception:
        kr_holidays = set()

    def _sff(v):
        try: return float(v)
        except Exception:
            import re as _re
            nums = _re.findall(r"[\d.]+", str(v))
            return float(nums[0]) if nums else 0.0

    # 멤버 팀 코드
    member_team = next((k for k, v in members.items() if v == name), None)

    def _shift_for_date(d):
        idx = (d - BASE_DATE).days % 20
        s1, s2, s3, off = CYCLE_20[idx]
        return {"1근_조": s1, "2근_조": s2, "3근_조": s3, "휴무_조": off,
                "1근_근무자": members.get(s1,""), "2근_근무자": members.get(s2,""),
                "3근_근무자": members.get(s3,""), "휴무_근무자": members.get(off,"")}

    def _apply_leaves(sh, d):
        result = dict(sh)
        result["is_2person"] = False
        for lv in base_leaves:
            try:
                ls = _dt.date.fromisoformat(lv["start"][:10])
                le = _dt.date.fromisoformat(lv["end"][:10])
            except Exception:
                continue
            if ls <= d <= le:
                absent = lv["name"]
                if result["1근_근무자"] == absent:
                    result["is_2person"] = True
                    result["주간_근무자"] = result["2근_근무자"]
                    result["야간_근무자"] = result["3근_근무자"]
                    result["주간_조"] = result["2근_조"]
                    result["야간_조"] = result["3근_조"]
                    result["leave_person"] = absent
                    result["leave_type"] = lv["type"]
                elif result["2근_근무자"] == absent:
                    result["is_2person"] = True
                    result["주간_근무자"] = result["1근_근무자"]
                    result["야간_근무자"] = result["3근_근무자"]
                    result["주간_조"] = result["1근_조"]
                    result["야간_조"] = result["3근_조"]
                    result["leave_person"] = absent
                    result["leave_type"] = lv["type"]
                elif result["3근_근무자"] == absent:
                    result["is_2person"] = True
                    result["주간_근무자"] = result["1근_근무자"]
                    result["야간_근무자"] = result["2근_근무자"]
                    result["주간_조"] = result["1근_조"]
                    result["야간_조"] = result["2근_조"]
                    result["leave_person"] = absent
                    result["leave_type"] = lv["type"]
                break
        return result

    days_in_month = _cal.monthrange(year, month)[1]
    salary_rows = []

    for fd in range(1, days_in_month + 1):
        fd_date = _dt.date(year, month, fd)
        fds = fd_date.strftime("%Y-%m-%d")
        fhol = fd_date in kr_holidays
        frow = {c: 0.0 for c in SCOLS}
        frow["날짜"] = f"{month:02d}/{fd:02d}"

        sh2 = None
        if fds in daily_details and daily_details[fds].get("shift"):
            sh2 = daily_details[fds]["shift"]
        else:
            auto = _shift_for_date(fd_date)
            sh2 = _apply_leaves(auto, fd_date)

        if not sh2:
            continue

        # 본인이 직접 휴가 등록된 날은 is_2person 여부와 무관하게 휴가비근로로 처리
        person_leave_type = None
        for lv in base_leaves:
            try:
                ls = _dt.date.fromisoformat(lv["start"][:10])
                le = _dt.date.fromisoformat(lv["end"][:10])
            except Exception:
                continue
            if ls <= fd_date <= le and lv["name"] == name:
                person_leave_type = lv["type"]
                break
        if person_leave_type:
            if person_leave_type == "공가":
                frow["공가"] = 8.0
            else:
                frow["휴가비근로"] = 8.0
            frow["일별합계"] = frow["공가"] + frow["휴가비근로"]
            salary_rows.append(frow)
            continue

        if sh2.get("is_2person"):
            lp = sh2.get("leave_person","") or sh2.get("3근_근무자","")
            dw = sh2.get("주간_근무자","") or sh2.get("1근_근무자","")
            nw = sh2.get("야간_근무자","") or sh2.get("2근_근무자","")
            lv_type = sh2.get("leave_type","") or sh2.get("3근_비고","")
            nb = NIGHT_HOURS.get("야간", 7.5)
            dot = _sff(sh2.get("1근_연장", 4) or 4)
            not_ = _sff(sh2.get("2근_연장", 4) or 4)
            if name == lp:
                if lv_type == "공가": frow["공가"] = 8.0
                else: frow["휴가비근로"] = 8.0
            elif name == dw:
                if fhol: frow["유휴근로"]=8.0; frow["휴일연장"]=dot; frow["휴일비근로"]=8.0
                else: frow["정상근로"]=8.0; frow["연장근로"]=dot
            elif name == nw:
                if fhol: frow["유휴근로"]=8.0; frow["야간근로"]=nb; frow["휴일연장"]=not_; frow["휴일비근로"]=8.0
                else: frow["정상근로"]=8.0; frow["야간근로"]=nb; frow["연장근로"]=not_
        else:
            for sk, 근k in [("1근_근무자","1근"),("2근_근무자","2근"),("3근_근무자","3근")]:
                if sh2.get(sk) == name:
                    nb3 = NIGHT_HOURS.get(근k, 0)
                    ot3 = _sff(sh2.get(f"{근k}_연장", 0) or 0)
                    dy3 = _sff(sh2.get(f"{근k}_주간연장") or ot3)
                    ny3 = _sff(sh2.get(f"{근k}_야간연장") or 0)
                    if fhol:
                        frow["유휴근로"]=8.0; frow["휴일비근로"]=8.0
                        frow["휴일연장"]=dy3+ny3
                        if nb3 > 0: frow["야간근로"] = nb3
                    else:
                        frow["정상근로"]=8.0; frow["연장근로"]=dy3
                        if nb3 > 0: frow["야간근로"] = nb3+ny3
                    break

        ft = sum(frow[c] for c in SCOLS)
        frow["일별합계"] = ft
        if ft > 0:
            salary_rows.append(frow)

    # 교대주기별 연장 계산
    ms = _dt.date(year, month, 1)
    next_mo = month % 12 + 1
    next_yr = year + (1 if month == 12 else 0)
    me = _dt.date(next_yr, next_mo, 1) - _dt.timedelta(days=1)
    ss3 = ms - _dt.timedelta(days=6)
    se3 = me + _dt.timedelta(days=6)

    # Source 1: 이미 계산된 salary_rows에서 연장 overtime 추출 (대근 포함, 가장 정확)
    obd: dict = {}
    for row in salary_rows:
        parts = row["날짜"].split("/")  # "MM/DD"
        date_key = f"{year}-{parts[0]}-{parts[1]}"
        ot = (row.get("연장근로") or 0) + (row.get("휴일연장") or 0)
        if ot > 0:
            obd[date_key] = ot

    # Source 2: 월 경계 전후 6일치 (다른 월의 cycle block 포함용)
    ext_prev = load_daily_detail_month(
        ss3.year, ss3.month
    ) if ss3.month != month or ss3.year != year else {}
    ext_next = load_daily_detail_month(
        se3.year, se3.month
    ) if se3.month != month or se3.year != year else {}
    ext_details = {**ext_prev, **ext_next}

    for ds5, dd5 in ext_details.items():
        d5 = _dt.date.fromisoformat(ds5)
        if not (ss3 <= d5 < ms or me < d5 <= se3):
            continue
        sh5 = dd5.get("shift", {}) if dd5 else {}
        if sh5 and not sh5.get("is_2person"):
            for sk5, ok5 in [("1근_근무자","1근"),("2근_근무자","2근"),("3근_근무자","3근")]:
                if sh5.get(sk5) == name:
                    ot5 = _sff(sh5.get(f"{ok5}_연장", 0) or 0)
                    if ot5 > 0: obd[ds5] = ot5
                    break
        elif sh5 and sh5.get("is_2person"):
            # 2인 근무: 주간/야간 연장 합산
            dw5 = sh5.get("주간_근무자","") or sh5.get("1근_근무자","")
            nw5 = sh5.get("야간_근무자","") or sh5.get("2근_근무자","")
            if name == dw5:
                ot5 = _sff(sh5.get("1근_연장", 4) or 4)
                if ot5 > 0: obd[ds5] = ot5
            elif name == nw5:
                ot5 = _sff(sh5.get("2근_연장", 4) or 4)
                if ot5 > 0: obd[ds5] = ot5

    # 근무 시퀀스
    wseq = []
    cur = ss3
    while cur <= se3:
        idx = (cur - BASE_DATE).days % 20
        c1, c2, c3, _ = CYCLE_20[idx]
        if member_team and member_team in (c1, c2, c3):
            wseq.append((cur, obd.get(cur.strftime("%Y-%m-%d"), 0)))
        cur += _dt.timedelta(days=1)

    # 연속 블록으로 그룹화
    blks = []
    if wseq:
        cb = [wseq[0]]
        for i in range(1, len(wseq)):
            if (wseq[i][0] - wseq[i-1][0]).days == 1:
                cb.append(wseq[i])
            else:
                blks.append(cb); cb = [wseq[i]]
        blks.append(cb)

    mblks = [b for b in blks if b[-1][0] >= ms and b[0][0] <= me]

    cycle_blocks = []
    for b in mblks:
        total_ot = int(sum(x[1] for x in b))
        cycle_blocks.append({
            "start": b[0][0].strftime("%m/%d"),
            "end": b[-1][0].strftime("%m/%d"),
            "total_ot": total_ot,
            "remaining": max(0, 12 - total_ot),
            "exceeded": total_ot >= 12,
            "warning": total_ot >= 10,
        })

    # 월합계
    totals = {c: sum(r[c] for r in salary_rows) for c in SCOLS}
    totals["일별합계"] = sum(r["일별합계"] for r in salary_rows)

    return {
        "salary_rows": salary_rows,
        "totals": totals,
        "cycle_blocks": cycle_blocks,
        "scols": SCOLS,
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
