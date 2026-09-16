"""
SYSTEM_GUIDE.md -> DOCX + PDF 변환 스크립트
"""
import re
from pathlib import Path

MD_PATH = Path(__file__).parent / "SYSTEM_GUIDE.md"
DOCX_PATH = Path(__file__).parent / "KG업무도우미_시스템가이드.docx"
PDF_PATH  = Path(__file__).parent / "KG업무도우미_시스템가이드.pdf"

text = MD_PATH.read_text(encoding="utf-8")

# ──────────────────────────────
# DOCX 생성
# ──────────────────────────────
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 기본 폰트
style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10)

def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.name = "맑은 고딕"
    return p

def add_table_from_md(doc, lines):
    rows = [l for l in lines if l.startswith("|") and "---" not in l]
    if not rows:
        return
    cells_list = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    col_count = max(len(c) for c in cells_list)
    t = doc.add_table(rows=len(cells_list), cols=col_count)
    t.style = "Table Grid"
    for ri, cells in enumerate(cells_list):
        for ci, val in enumerate(cells):
            cell = t.cell(ri, ci)
            cell.text = val
            run = cell.paragraphs[0].runs
            if run and ri == 0:
                run[0].bold = True
            if run:
                run[0].font.name = "맑은 고딕"
                run[0].font.size = Pt(9)

i = 0
lines = text.splitlines()
while i < len(lines):
    line = lines[i]

    # 제목 처리
    if line.startswith("### "):
        add_heading(doc, line[4:], 3)
    elif line.startswith("## "):
        add_heading(doc, line[3:], 2)
    elif line.startswith("# "):
        add_heading(doc, line[2:], 1)

    # 표 처리
    elif line.startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].startswith("|"):
            table_lines.append(lines[i])
            i += 1
        add_table_from_md(doc, table_lines)
        continue

    # 코드 블록
    elif line.startswith("```"):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith("```"):
            code_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Courier New"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 수평선
    elif line.strip() == "---":
        doc.add_paragraph("─" * 50)

    # 빈 줄
    elif line.strip() == "":
        pass

    # 일반 텍스트 / 목록
    else:
        # 굵게 처리(**text**)
        clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        # 인라인 코드(`code`)
        clean = re.sub(r"`(.+?)`", r"\1", clean)
        # 링크 [text](url) → text (url)
        clean = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", clean)

        if line.startswith("- "):
            p = doc.add_paragraph(clean[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(re.sub(r"^\d+\. ", "", clean), style="List Number")
        else:
            p = doc.add_paragraph(clean)

        for run in p.runs:
            run.font.name = "맑은 고딕"
            run.font.size = Pt(10)

    i += 1

doc.save(DOCX_PATH)
print(f"DOCX 저장: {DOCX_PATH}")

# ──────────────────────────────
# PDF 생성 (fpdf2 + 맑은 고딕)
# ──────────────────────────────
from fpdf import FPDF

FONT_PATH = r"C:\Windows\Fonts\malgunbd.ttf"   # 맑은 고딕 Bold
FONT_REG  = r"C:\Windows\Fonts\malgun.ttf"     # 맑은 고딕 Regular
FONT_MONO = r"C:\Windows\Fonts\cour.ttf"       # Courier New

class PDF(FPDF):
    def header(self):
        self.set_font("malgun", size=8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 6, "KG스틸 업무도우미 시스템 가이드", align="R")
        self.ln(4)

    def footer(self):
        self.set_y(-12)
        self.set_font("malgun", size=8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 6, f"- {self.page_no()} -", align="C")

pdf = PDF()
pdf.add_font("malgun",  fname=FONT_REG,  uni=True)
pdf.add_font("malgunB", fname=FONT_PATH, uni=True)
pdf.add_font("mono",    fname=FONT_MONO, uni=True)
pdf.set_auto_page_break(auto=True, margin=15)
pdf.add_page()

i = 0
lines = text.splitlines()
while i < len(lines):
    line = lines[i]

    if line.startswith("# ") and not line.startswith("## "):
        pdf.set_font("malgunB", size=18)
        pdf.set_text_color(26, 26, 46)
        pdf.multi_cell(0, 10, line[2:])
        pdf.ln(2)

    elif line.startswith("## "):
        pdf.ln(4)
        pdf.set_font("malgunB", size=14)
        pdf.set_text_color(30, 100, 180)
        pdf.multi_cell(0, 8, line[3:])
        pdf.set_draw_color(30, 100, 180)
        pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 170, pdf.get_y())
        pdf.ln(2)

    elif line.startswith("### "):
        pdf.ln(2)
        pdf.set_font("malgunB", size=11)
        pdf.set_text_color(50, 50, 50)
        pdf.multi_cell(0, 7, line[4:])

    elif line.startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].startswith("|"):
            if "---" not in lines[i]:
                table_lines.append([c.strip() for c in lines[i].strip("|").split("|")])
            i += 1

        if table_lines:
            col_n = len(table_lines[0])
            col_w = 170 / col_n
            pdf.ln(2)
            for ri, row in enumerate(table_lines):
                if ri == 0:
                    pdf.set_font("malgunB", size=8.5)
                    pdf.set_fill_color(200, 220, 255)
                    fill = True
                else:
                    pdf.set_font("malgun", size=8.5)
                    pdf.set_fill_color(245, 245, 245)
                    fill = (ri % 2 == 0)
                pdf.set_text_color(30, 30, 30)
                for ci, cell in enumerate(row[:col_n]):
                    pdf.cell(col_w, 7, cell[:40], border=1, fill=fill)
                pdf.ln()
            pdf.ln(2)
        continue

    elif line.startswith("```"):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith("```"):
            code_lines.append(lines[i])
            i += 1
        pdf.set_font("mono", size=7.5)
        pdf.set_fill_color(240, 240, 240)
        pdf.set_text_color(50, 50, 50)
        for cl in code_lines:
            pdf.set_x(20)
            pdf.cell(0, 5, cl[:100], fill=True)
            pdf.ln()
        pdf.ln(2)

    elif line.strip() == "---":
        pdf.set_draw_color(180, 180, 180)
        pdf.line(10, pdf.get_y() + 2, 200, pdf.get_y() + 2)
        pdf.ln(4)

    elif line.strip() == "":
        pdf.ln(3)

    else:
        clean = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        clean = re.sub(r"`(.+?)`", r"\1", clean)
        clean = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", clean)

        pdf.set_text_color(40, 40, 40)
        if line.startswith("- "):
            pdf.set_font("malgun", size=10)
            pdf.multi_cell(170, 6, "  • " + clean[2:])
        elif re.match(r"^\d+\. ", line):
            pdf.set_font("malgun", size=10)
            pdf.multi_cell(170, 6, "  " + clean)
        elif line.startswith("> "):
            pdf.set_font("malgun", size=9)
            pdf.set_text_color(100, 100, 100)
            pdf.multi_cell(170, 6, "  " + clean[2:])
        else:
            pdf.set_font("malgun", size=10)
            pdf.multi_cell(170, 6, clean)

    i += 1

pdf.output(str(PDF_PATH))
print(f"PDF  저장: {PDF_PATH}")
print("완료!")
