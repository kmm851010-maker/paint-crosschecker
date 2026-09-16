"""
KG Steel 업무도우미 & KG OPS 매뉴얼 생성 스크립트
실행: python docs/generate_docs.py
출력: docs/web_manual.html, docs/web_manual.docx
      docs/app_manual.html, docs/app_manual.docx
"""

import base64
import re
from pathlib import Path
from io import BytesIO
from PIL import Image

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT    = Path(__file__).parent
WEB_SS  = ROOT / "screenshots" / "web"
APP_SS  = ROOT / "screenshots" / "app"

# ─────────────────────────────────────────
# 문서 콘텐츠 정의
# ─────────────────────────────────────────

WEB_DOC = {
    "title": "업무도우미 사용 매뉴얼",
    "subtitle": "KG스틸 당진생산지원팀 전용 업무 지원 시스템",
    "color": "#4B2D8E",
    "sections": [
        {
            "heading": "1. 시스템 개요",
            "sub": [],
            "content": [
                ("text", "업무도우미는 KG스틸 당진생산지원팀 사무실에서 사용하는 웹 기반 업무 지원 플랫폼입니다. 생산계획서 분석, ERP 교차검증, 일일 작업일지 작성, 근무표 관리, 창고 재고 현황 조회 등 팀의 반복 업무를 자동화하여 업무 효율을 크게 높입니다."),
                ("heading3", "도입 효과"),
                ("bullets", [
                    "생산계획서 → ERP 데이터 교차검증 자동화로 입력 오류 즉시 감지",
                    "휴가 신청 한 번으로 근무표·작업일지·근무통계 전체 자동 반영 (수동 입력 불필요)",
                    "일일 작업일지 작성 및 월 통합 엑셀 자동 생성, 이메일 발송 원클릭 처리",
                    "근무표·근무통계 자동 집계로 수기 정리 불필요",
                    "창고 드럼 위치를 실시간으로 확인 (KG OPS 앱 연동)",
                    "반품 드럼 종류별(기술·불량·무상) 현황 추적",
                ]),
                ("img", "01_main.png", "업무도우미 메인 화면"),
            ],
        },
        {
            "heading": "2. 접속 방법",
            "content": [
                ("text", "PC 웹 브라우저(Chrome 권장)에서 아래 주소로 접속합니다. 아이디와 비밀번호를 입력한 후 '로그인' 버튼을 클릭하세요."),
                ("info", "접속 주소는 팀 내부 공지를 통해 확인하세요."),
                ("img", "00_login.png", "로그인 화면"),
                ("bullets", [
                    "아이디: 부여받은 개인 아이디 입력",
                    "비밀번호: 팀 공용 비밀번호 입력",
                    "로그인 후 왼쪽 사이드바에서 메뉴를 선택합니다",
                ]),
            ],
        },
        {
            "heading": "3. 근무표",
            "content": [
                ("text", "월별 근무 일정을 달력 형태로 확인합니다. 4조3교대 로테이션 기준으로 각 직원의 1근(06:30~14:30) · 2근(14:30~22:30) · 3근(22:30~06:30) · 휴무 일정을 자동 계산하여 표시합니다. 휴가를 등록하면 해당 날짜의 작업일지와 근무 통계에도 자동으로 반영됩니다."),
                ("img", "02_근무표.png", "근무표 화면"),
                ("heading3", "근무표 조회 방법"),
                ("numbered", [
                    "왼쪽 사이드바에서 [근무표] 클릭",
                    "이름(또는 조) · 연도 · 월 선택",
                    "달력에서 근무 일정 확인 (1근/2근/3근/휴무/휴가 배지로 표시)",
                    "달력의 날짜 칸을 클릭하면 해당 일 전체 인원의 근무 상세 카드 표시",
                    "월 이동: 상단 ◀ ▶ 버튼으로 이전/다음 달 조회",
                ]),
                ("heading3", "📝 휴가 신청 방법"),
                ("text", "근무표 상단의 [📝 휴가신청서 작성] 버튼을 클릭하면 입력 창이 펼쳐집니다. 아래 항목을 작성 후 [✅ 등록] 버튼을 누르면 즉시 저장됩니다."),
                ("numbered", [
                    "[근무표] 메뉴 접속 → 상단 [📝 휴가신청서 작성] 클릭하여 펼치기",
                    "대상자: 드롭다운에서 휴가 대상 직원 선택",
                    "구분: 해당 휴가 유형 선택 — 정기휴가 / 연차 / 특별휴가 / 명휴 / 생일휴가 / 공가 / 공상휴업 / 산재 / 휴직 / 대휴 / 교육 / 결근 / 조퇴 / 외출 / 청원휴가 / 공휴",
                    "시작일 / 종료일: 달력에서 해당 날짜 선택 (하루짜리면 같은 날로 설정)",
                    "대근자: 해당 날 대신 들어오는 직원이 있으면 선택, 없으면 공란",
                    "[✅ 등록] 버튼 클릭 → 즉시 저장 및 전체 시스템 자동 반영",
                ]),
                ("info", "공휴(공휴일 근무) 선택 시, 선택 기간 중 법정공휴일이 아닌 날이 포함되면 경고 메시지가 표시됩니다. 내용을 확인하고 체크박스에 체크해야 등록이 가능합니다."),
                ("heading3", "등록된 일정 조회 및 삭제"),
                ("bullets", [
                    "휴가신청서 작성 창 하단에서 연도·월 필터로 해당 기간 일정 목록 조회",
                    "직원명 | 구분 | 기간 | 대근자 형태로 목록 표시",
                    "각 항목 우측 [삭제] 버튼 클릭으로 즉시 삭제",
                    "삭제 시 해당 기간 중 오늘 이후 날짜의 연관 작업일지 자동 초기화 (과거 확정 일지는 보존)",
                ]),
                ("heading3", "⚡ 휴가 등록 후 자동 연동 효과"),
                ("text", "휴가를 등록하면 추가 작업 없이 아래 항목이 전체 시스템에 자동으로 반영됩니다. 근무자를 일일이 바꾸거나 작업일지를 다시 설정할 필요가 없습니다."),
                ("bullets", [
                    "📅 근무표: 해당 날짜에 해당 직원이 휴가(갈색 배지)로 자동 표시",
                    "📋 작업일지: 해당 날짜 선택 시 3근→ 2인 근무 체계(주간 06:30~18:30 / 야간 18:30~06:30)로 자동 전환",
                    "👥 인원 현황 자동 배정: 1근·2근·3근 중 누가 빠지는지 계산하여 나머지 2인이 주간/야간에 자동 배치 (수동 입력 불필요)",
                    "📝 비고 자동 입력: 일반 휴가 → '대휴 연장4H', 공휴 → '공휴일 휴일연장대근'으로 자동 채움",
                    "📊 근무 통계: 휴가자 휴가일수 · 대근자 대근횟수 자동 집계",
                ]),
            ],
        },
        {
            "heading": "4. 근무 통계",
            "content": [
                ("text", "직원별 월간 근무 시간, 야간 근무 시간, 휴가 사용 현황, 대근 내역 등을 자동으로 집계합니다. 저장된 작업일지 데이터와 휴가 신청 내역을 바탕으로 계산하므로 수기 정리가 불필요합니다."),
                ("img", "03_근무_통계.png", "근무 통계 화면"),
                ("heading3", "화면 구성 및 주요 기능"),
                ("bullets", [
                    "조회 대상자 · 연도 · 월 선택 → 해당 직원의 월 통계 즉시 표시",
                    "직원별 월간 근무 일수 · 기본 근로시간 자동 집계",
                    "대근 연장근로 시간 계산 (휴가자 대신 2인 근무 시 대근 1회당 +4H 자동 반영)",
                    "야간근로 시간 구분 계산 (22:00~06:00 기준)",
                    "휴가 사용 일수 및 내역 자동 집계 (등록된 휴가 신청 기반)",
                    "대근 횟수 및 대근 내역 날짜·대상자별 자동 기록",
                    "저장된 작업일지 없는 날은 교대 로테이션 기준 예측값으로 보완",
                    "엑셀 파일 다운로드 지원",
                ]),
                ("heading3", "통계 항목 설명"),
                ("bullets", [
                    "근무일수: 1근/2근/3근으로 실제 출근한 날 수",
                    "기본근로: 근무일수 × 8시간",
                    "대근 연장: 타인 휴가 시 2인 근무로 추가된 연장 시간 합계",
                    "야간근로: 22:00~06:00 사이에 해당하는 근무 시간 합계",
                    "휴가일수: 연차·정기휴가 등 등록된 휴가 일수 합계",
                    "대근횟수: 타인의 휴가를 대신 근무한 총 횟수",
                ]),
                ("heading3", "대근 내역 (예: 4회 · 계 16H)"),
                ("text", "대근 내역 항목은 해당 직원이 타인 휴가 시 2인 근무 체계로 대신 근무한 횟수와 총 시간이 자동 집계됩니다. 상세 내역을 펼치면 날짜별로 확인할 수 있습니다."),
                ("bullets", [
                    "날짜: 대근 발생 날짜",
                    "휴가자: 휴가를 쓴 직원 이름",
                    "휴가구분: 정기휴가·연차·공휴 등 휴가 유형",
                    "시간: 해당 날 대근 연장 시간 (1회 = 4H 기본)",
                    "구분: 주간 대근 또는 야간 대근",
                ]),
                ("heading3", "휴가 내역 (월별 / 연간)"),
                ("text", "화면 하단에서 해당 월 휴가 내역과 해당 연도 전체 휴가 내역을 별도로 확인할 수 있습니다."),
                ("bullets", [
                    "8월 휴가 내역 (N일): 해당 월에 사용한 휴가 일자와 구분(연차/정기휴가 등) 목록",
                    "2026년 전체 휴가 내역 (N일): 해당 연도 누적 휴가 일수 및 날짜별 내역",
                ]),
                ("heading3", "급여시간표 자세히 보기"),
                ("text", "[급여시간표 자세히 보기] 버튼을 클릭하면 해당 월의 일별 근무 상세 내역을 표로 확인할 수 있습니다."),
                ("bullets", [
                    "일자별 근무 구분 (1근/2근/3근/휴무/휴가)",
                    "기본근로·연장근로·야간근로·대근연장 시간 일별 표시",
                    "급여 산정에 필요한 시간 데이터를 날짜별로 일괄 확인",
                ]),
                ("heading3", "교대주기별 연장 시간"),
                ("text", "교대주기별 연장 시간 항목은 3근(야간 → 주간) 교대 구간에서 발생하는 연장 시간을 교대 주기별로 집계합니다."),
                ("bullets", [
                    "4조3교대는 20일 주기로 로테이션하며, 주기가 바뀌는 시점에 연장 근무가 발생할 수 있음",
                    "교대 전환 연장(야간 연속 근무 등)이 있는 주기를 별도로 구분하여 확인 가능",
                    "급여 정산 및 근로시간 검토 시 활용",
                ]),
                ("info", "근무 통계는 저장된 작업일지와 휴가 신청 데이터를 자동 연동합니다. 작업일지를 저장하지 않은 날은 교대 로테이션 기준 예측값이 사용됩니다. 정확도를 높이려면 매일 작업일지를 저장하세요."),
            ],
        },
        {
            "heading": "5. 일일 작업 일지",
            "content": [
                ("text", "당일 작업 내용, 인원 현황, 안전 관리 사항, 특이 사항을 기록합니다. 근무표에 휴가가 등록된 날짜를 선택하면 2인 근무 체계로 자동 전환되고, 근무자 이름이 자동으로 채워져 별도로 수정할 필요가 없습니다."),
                ("img", "04_일일_작업_일지.png", "일일 작업 일지 화면"),
                ("heading3", "① 작업 일자 선택"),
                ("bullets", [
                    "상단 날짜 선택기에서 작업 일자 클릭 (기본값: 오늘 날짜, 06:30 이전이면 전일로 자동 설정)",
                    "이전에 저장된 데이터가 있으면 날짜 변경 즉시 자동 불러오기",
                    "과거 날짜 선택 → 이전 작업일지 수정 가능",
                ]),
                ("heading3", "② 인원 현황 (휴가 자동 연동)"),
                ("text", "날짜 선택 시 4조3교대 로테이션 기준으로 당일 근무자가 자동으로 표시됩니다. 휴가가 등록된 날짜는 별도 안내 배너와 함께 2인 근무 체계로 자동 전환됩니다."),
                ("bullets", [
                    "일반 3교대일 때: 1근(06:30~14:30) / 2근(14:30~22:30) / 3근(22:30~06:30) / 휴무 자동 표시",
                    "휴가 등록된 날짜: 주간(06:30~18:30) / 야간(18:30~06:30) 2인 체계로 자동 전환, 근무자 이름 자동 배치",
                    "근무자 이름은 자동 입력되어 있으므로 수정이 필요한 경우에만 변경",
                    "연장근무 있을 시: 연장 시작 시간 · 종료 시간 선택 → 주간연장·야간연장 시간 자동 계산",
                    "휴무 구분: 드롭다운에서 교대휴무 / 주휴휴무 / 연차 / 정기휴가 등 선택",
                ]),
                ("heading3", "③ 업무 현황 입력"),
                ("bullets", [
                    "항목별(페인트 하차·공급, 신나 하차·공급·크롬 공급, 공드럼, 페보루, 페신너, 반품·불량, 코터롤, 필름, AGV) 근별 수량 입력",
                    "숫자 직접 입력 또는 덧셈 수식(예: 10+5) 입력 지원",
                    "일합계: 각 행의 입력값을 자동 합산하여 표시",
                    "월누계: Google Sheets에서 해당 월 전체(오늘 제외) 누적 수량 자동 조회",
                ]),
                ("heading3", "④ 안전 관리 사항"),
                ("bullets", [
                    "작업 절차 준수 여부, 안전장치 점검, 급출발/급정거 금지 등 6개 항목 체크",
                    "각 항목을 근별(1근/2근/3근, 또는 주간/야간)로 체크",
                    "기본값: 전체 준수(체크) 상태로 설정됨",
                ]),
                ("heading3", "⑤ 특이 사항"),
                ("bullets", [
                    "당일 공지사항, 사고/이상 내용, 전달 사항 등 자유 텍스트 입력",
                ]),
                ("heading3", "저장 · 다운로드 · 메일 발송"),
                ("numbered", [
                    "[일일 작업 일지] 메뉴 클릭",
                    "작업 일자 선택 → 기존 저장 데이터 있으면 자동 불러오기",
                    "인원 현황 확인 (휴가 있으면 2인 근무 자동 전환 안내 배너 표시)",
                    "업무 현황 항목별 수량 입력",
                    "안전 관리 사항 체크 확인 및 특이 사항 입력",
                    "[💾 전체 저장 (Google Sheets)] 버튼 클릭 → 저장 완료",
                    "저장 후 당월 통합 Excel 자동 생성 → [📥 다운로드] 버튼으로 다운로드",
                    "필요 시 [✉️ 통합일지 메일 전송] 버튼 클릭 → 수신자·제목·본문 입력 후 발송",
                ]),
                ("info", "저장 완료 후 20초 이내 재저장 방지 쿨다운이 적용됩니다. 저장된 데이터를 무시하고 새로 작성하려면 [🆕 새로 작성] 버튼을 사용하세요. 당일 06:30 이전 작업은 전일 일지로 자동 분류됩니다."),
            ],
        },
        {
            "heading": "6. 입고 관리 (생산계획서 교차검증)",
            "content": [
                ("text", "생산계획서(이미지 또는 엑셀)와 ERP 입고명세서를 비교하여 입고 수량과 품목의 일치 여부를 자동으로 검증합니다. 수기 대조 작업을 완전히 대체합니다."),
                ("img", "06_입고_관리.png", "입고 관리 화면"),
                ("heading3", "검증 절차"),
                ("numbered", [
                    "생산계획서(이미지 또는 엑셀 파일)를 업로드 → AI가 자동으로 입고 예정 품목 추출",
                    "추출된 리스트 확인 후 실제 입고 작업 진행",
                    "ERP에서 입고명세서를 캡처하거나 다운로드",
                    "ERP 입고명세서 업로드 → 자동 교차검증 실행",
                    "일치/부족/초과 항목이 색상으로 구분되어 표시",
                    "결과를 엑셀 파일로 저장 가능",
                ]),
                ("heading3", "검증 결과 색상 의미"),
                ("bullets", [
                    "✅ 일치(초록): 계획과 실제 입고 수량 일치",
                    "⚠️ 초과(노랑): 계획보다 많이 입고됨",
                    "❌ 부족(빨강): 계획보다 적게 입고됨",
                    "⬜ 미입고(회색): 계획은 있으나 입고되지 않음",
                ]),
            ],
        },
        {
            "heading": "7. 재고 현황 (웹)",
            "content": [
                ("text", "KG OPS 앱으로 등록된 드럼 재고를 PC에서 조회합니다. 섹터별 현황과 날짜별 이동 이력을 확인할 수 있습니다."),
                ("img", "05_재고_현황.png", "재고 현황 - 섹터별"),
                ("img", "05_재고_현황_날짜별.png", "재고 현황 - 날짜별 이력"),
                ("heading3", "탭 구성"),
                ("bullets", [
                    "섹터별 현황: 각 보관 구역의 드럼 목록 조회",
                    "날짜별 이력: 특정 기간 내 신규등록·이동·라인입고·반품 이력 조회",
                ]),
                ("info", "날짜별 이력은 06:30 기준으로 하루가 구분됩니다. (예: 8월 29일 = 29일 06:30 ~ 30일 06:30)"),
            ],
        },
        {
            "heading": "8. 반품 관리",
            "content": [
                ("text", "기술 반품, 불량 반품, 무상 반품 드럼을 종류별로 관리합니다. 반품 사유와 처리 현황을 기록하고 추적할 수 있습니다."),
                ("img", "07_반품_관리.png", "반품 관리 화면"),
                ("heading3", "반품 유형"),
                ("bullets", [
                    "🔴 불량 반품: 품질 불량으로 인한 반품",
                    "🟡 기술 반품: 기술적 사유로 인한 반품",
                    "🔵 무상 반품: 무상 교환 처리 반품",
                ]),
                ("heading3", "처리 방법"),
                ("numbered", [
                    "[반품 관리] 메뉴 클릭",
                    "상단 필터로 반품 유형 선택",
                    "드럼 LOT 번호로 해당 드럼 검색",
                    "반품 상태 업데이트 및 처리 완료 등록",
                ]),
            ],
        },
        {
            "heading": "9. 주의사항 및 자주 묻는 질문",
            "content": [
                ("heading3", "주의사항"),
                ("bullets", [
                    "작업일지 저장 후 20초 이내에는 재저장이 제한됩니다. 쿨다운 후 재시도하세요.",
                    "작업일지를 빈칸 상태에서 실수로 저장하면 Google Sheets 데이터가 덮어씌워집니다. 저장 전 입력값을 반드시 확인하세요.",
                    "휴가를 삭제하면 오늘 이후 날짜의 관련 작업일지가 자동 초기화됩니다. 과거 확정 일지는 보존됩니다.",
                    "한 날짜에 2명 이상 동시 휴가 등록은 지원하지 않습니다. 한 명씩 등록하세요.",
                    "생산계획서 교차검증 시 이미지 화질이 낮으면 AI 인식 정확도가 떨어질 수 있습니다. 선명한 이미지를 사용하세요.",
                    "재고 현황은 KG OPS 앱에서 등록한 데이터를 실시간으로 불러옵니다. 앱에서 등록하지 않은 드럼은 표시되지 않습니다.",
                    "인터넷 연결이 불안정하면 저장이 실패할 수 있습니다. 저장 실패 메시지 확인 후 재시도하세요.",
                ]),
                ("heading3", "자주 묻는 질문"),
                ("bullets", [
                    "Q. 작업일지에 저장했는데 다운받은 파일에 내용이 없어요 → 저장 성공 메시지를 확인했는지 확인하세요. 실패 시 오류 메시지가 표시됩니다.",
                    "Q. 메뉴를 이동했다 돌아오니 작업일지 입력값이 사라졌어요 → 이전에 저장한 데이터가 있으면 날짜 선택 시 자동으로 불러옵니다. '새로 작성' 버튼을 클릭한 경우 입력값이 초기화됩니다.",
                    "Q. 휴가를 등록했는데 작업일지에 2인 근무로 바뀌지 않아요 → 작업일지 메뉴에서 해당 날짜를 다시 선택하면 자동으로 반영됩니다.",
                    "Q. 월누계 숫자가 틀린 것 같아요 → 월누계는 당일을 제외한 해당 월 전체 합산입니다. 저장 후 다음 날부터 오늘 데이터가 포함됩니다.",
                    "Q. 저장 버튼이 비활성화되어 있어요 → 저장 후 20초 쿨다운 중입니다. 남은 시간이 버튼 아래에 표시됩니다.",
                    "Q. 근무 통계 수치가 실제와 달라요 → 작업일지를 저장하지 않은 날은 교대 로테이션 예측값이 사용됩니다. 매일 작업일지를 저장하면 정확도가 높아집니다.",
                    "Q. 로그인이 안 돼요 → 아이디·비밀번호를 확인하고, 접속 주소가 맞는지 팀 내부 공지를 확인하세요.",
                    "Q. 화면이 계속 로딩 중이에요 → 브라우저를 새로고침(F5)하거나 Chrome 캐시를 지우고 다시 시도하세요.",
                ]),
                ("info", "위 내용으로 해결되지 않는 문제는 시스템 담당자에게 문의하세요."),
            ],
        },
    ],
}

APP_DOC = {
    "title": "KG OPS 사용 매뉴얼",
    "subtitle": "현장 드럼 재고 관리 모바일 앱",
    "color": "#4B2D8E",
    "sections": [
        {
            "heading": "1. 앱 소개",
            "content": [
                ("text", "KG OPS는 현장 작업자가 스마트폰으로 도료 드럼 재고를 실시간 관리하는 앱입니다. 드럼 바코드를 스캔하여 보관 위치(섹터)를 등록하고, 라인 입고 처리, 이력 조회까지 현장에서 즉시 처리할 수 있습니다."),
                ("heading3", "주요 기능"),
                ("bullets", [
                    "드럼 바코드 카메라 자동 스캔 및 LOT 번호 인식",
                    "섹터별 드럼 등록 및 이동 처리",
                    "라인 입고(재고 출고) 원터치 처리",
                    "현재 재고 섹터별 조회",
                    "날짜별 이동 이력 조회",
                ]),
                ("heading3", "도입 효과"),
                ("bullets", [
                    "종이 대장 작성 불필요 → 실시간 디지털 기록",
                    "드럼 위치를 즉시 확인 가능 → 찾는 시간 절감",
                    "입고·출고 이력 자동 기록 → 감사 추적 가능",
                    "사무실 PC(업무도우미)와 실시간 연동",
                ]),
            ],
        },
        {
            "heading": "2. 앱 설치 방법",
            "content": [
                ("text", "PC 업무도우미 사이드바 하단의 [⬇️ KG OPS 설치] 버튼을 클릭하여 APK 파일을 다운로드합니다."),
                ("heading3", "설치 절차"),
                ("numbered", [
                    "스마트폰으로 업무도우미 접속 → [KG OPS 설치] 버튼 클릭",
                    "APK 파일 다운로드 완료 후 파일 실행",
                    "'알 수 없는 앱' 설치 허용 (최초 1회)",
                    "설치 완료 후 KG OPS 앱 실행",
                ]),
                ("info", "Android 폰 전용입니다. 최초 설치 시 '알 수 없는 앱 설치 허용' 설정이 필요합니다."),
            ],
        },
        {
            "heading": "3. 드럼 스캔 및 섹터 등록",
            "content": [
                ("text", "카메라로 드럼 바코드를 스캔하면 LOT 번호와 품명, 제조사가 자동으로 인식됩니다. 인식된 드럼을 선택하여 보관 섹터에 등록합니다."),
                ("img", "01_main_scan.png", "메인 스캔 화면"),
                ("heading3", "스캔 방법"),
                ("numbered", [
                    "앱 실행 후 카메라 화면에 드럼 측면 바코드를 비춤 (스티커 바코드 또는 드럼 자체 코드)",
                    "바코드 인식 시 LOT 번호 · 품명 · 제조사가 자동으로 목록에 추가됨",
                    "여러 드럼을 연속으로 스캔 가능 (목록에 계속 추가)",
                    "잘못 인식된 항목은 해당 항목을 길게 눌러 삭제",
                ]),
                ("info", "바코드 인식이 잘 안 되는 경우: 조명이 밝은 곳에서 바코드에 카메라를 수평으로 맞추세요. 바코드가 찍히거나 오염된 경우 LOT 번호를 직접 검색하여 수동으로 추가할 수 있습니다."),
                ("img", "02_sector_select.png", "섹터 선택 화면"),
                ("heading3", "섹터 등록 방법"),
                ("numbered", [
                    "스캔된 드럼 목록에서 등록할 드럼의 체크박스 선택 (전체 선택 가능)",
                    "[섹터 등록] 버튼 클릭",
                    "드럼을 보관할 섹터 선택",
                    "[등록] 버튼으로 확정 → 서버에 즉시 저장",
                ]),
                ("heading3", "섹터 구성"),
                ("bullets", [
                    "입고존: 방금 입고된 드럼의 임시 보관 구역",
                    "신나자리: 신나류(희석제) 전용 보관 구역",
                    "0~3번자리 / 4~6번자리 / 7A~C자리 / 7D~Z자리: 외곽 창고 번호별 구역",
                    "8번자리 / 9번자리: 내부 창고 구역",
                    "반품자리: 반품 처리 대기 드럼 보관",
                    "창고주위: 기타 창고 주변 임시 보관 구역",
                ]),
                ("info", "이미 다른 섹터에 등록된 드럼을 스캔하면 이동 처리됩니다. 이전 섹터 기록은 이력에 자동 저장됩니다."),
            ],
        },
        {
            "heading": "4. 재고 현황 조회 및 라인입고",
            "content": [
                ("text", "현재 등록된 모든 드럼의 보관 위치를 섹터별로 조회합니다. 라인에 투입할 드럼은 선택 후 라인입고 처리하면 재고에서 자동으로 제거됩니다."),
                ("img", "03_inventory.png", "재고 현황 화면"),
                ("heading3", "재고 조회 방법"),
                ("numbered", [
                    "메인 화면에서 [재고현황] 버튼 클릭",
                    "[현재 재고] 탭 선택",
                    "섹터별로 드럼 목록 확인",
                    "특정 드럼은 LOT 번호로 검색 가능",
                ]),
                ("heading3", "라인입고 처리"),
                ("numbered", [
                    "재고 목록에서 라인에 투입할 드럼 체크박스 선택",
                    "[라인입고] 버튼 클릭",
                    "확인 후 처리 완료 → 재고에서 자동 제거",
                    "처리 내역은 날짜별 이력에서 확인 가능",
                ]),
                ("info", "라인입고 처리된 드럼은 재고 목록에서 제거되며, 이력에는 '라인입고'로 기록됩니다."),
            ],
        },
        {
            "heading": "5. 날짜별 이력 조회",
            "content": [
                ("text", "특정 기간 동안의 드럼 이동 이력을 조회합니다. 신규 등록, 섹터 이동, 라인입고, 반품완료 이력을 탭으로 구분하여 확인할 수 있습니다."),
                ("img", "04_history.png", "날짜별 이력 화면"),
                ("heading3", "조회 방법"),
                ("numbered", [
                    "[현재 재고] 탭 옆 [날짜별 이력] 탭 클릭",
                    "시작 날짜와 종료 날짜 입력 (YYYY-MM-DD 형식)",
                    "[조회] 버튼 클릭",
                    "신규등록 / 라인입고 / 반품완료 탭으로 구분 확인",
                ]),
                ("info", "날짜 기준은 06:30입니다. 예) 8월 29일 조회 = 29일 06:30 ~ 30일 06:30 사이 기록"),
                ("heading3", "이력 항목 설명"),
                ("bullets", [
                    "신규등록: 해당 드럼이 처음 시스템에 등록됨",
                    "이동: 보관 섹터가 변경됨",
                    "라인입고: 라인에 투입되어 재고에서 제거됨",
                    "반품완료: 반품 처리가 완료됨",
                ]),
            ],
        },
        {
            "heading": "6. 주의사항 및 자주 묻는 질문",
            "content": [
                ("heading3", "주의사항"),
                ("bullets", [
                    "라인입고 처리 후에는 재고 목록에서 즉시 제거됩니다. 잘못 처리한 경우 사무실 PC 업무도우미에서 이력을 확인하세요.",
                    "같은 LOT 번호를 두 번 스캔하면 섹터 이동으로 처리됩니다 (중복 등록 방지).",
                    "앱 사용 중 인터넷 연결이 끊기면 데이터 저장이 안 될 수 있습니다. 저장 완료 메시지를 반드시 확인하세요.",
                    "앱 업데이트가 있는 경우 업무도우미 사이드바 하단 [KG OPS 설치] 버튼으로 최신 APK를 재설치합니다.",
                ]),
                ("heading3", "자주 묻는 질문"),
                ("bullets", [
                    "Q. 바코드가 인식되지 않아요 → 바코드 위 이물질 제거 후 밝은 곳에서 다시 시도. 그래도 안 되면 LOT 번호 직접 입력",
                    "Q. 잘못된 섹터에 등록했어요 → 해당 드럼을 다시 스캔하여 올바른 섹터로 재등록하면 이동 처리됩니다",
                    "Q. 라인입고를 실수로 눌렀어요 → 사무실 PC 업무도우미의 재고 현황에서 이력 확인 후 담당자에게 수동 복원 요청",
                    "Q. 앱이 갑자기 꺼져요 → 스마트폰 저장공간 부족 또는 앱 권한 문제일 수 있습니다. 카메라 권한을 허용했는지 확인하세요",
                    "Q. 등록한 드럼이 재고 현황에 안 보여요 → 앱을 재시작하거나 [새로고침] 버튼을 눌러보세요",
                ]),
                ("info", "문제가 지속되면 사무실 PC 업무도우미에서 재고 현황과 이력을 먼저 확인한 후 담당자에게 문의하세요."),
            ],
        },
    ],
}


# ─────────────────────────────────────────
# HTML 생성
# ─────────────────────────────────────────

def img_to_b64(path: Path) -> str:
    if not path.exists():
        return ""
    with open(path, "rb") as f:
        return base64.b64encode(f.read()).decode()

def render_html(doc: dict, ss_dir: Path) -> str:
    color   = doc["color"]
    title   = doc["title"]
    sub     = doc["subtitle"]

    css = f"""
    * {{ margin: 0; padding: 0; box-sizing: border-box; }}
    body {{ font-family: 'Malgun Gothic', -apple-system, sans-serif; color: #1a1a1a; background: #f5f5f5; }}
    .cover {{ background: linear-gradient(135deg, {color} 0%, #2d1a5e 100%); color: #fff; padding: 60px 40px; }}
    .cover h1 {{ font-size: 32px; font-weight: 800; margin-bottom: 10px; }}
    .cover p  {{ font-size: 15px; opacity: 0.85; }}
    .toc {{ background: #fff; padding: 30px 40px; border-bottom: 3px solid {color}; }}
    .toc h2 {{ font-size: 16px; font-weight: 700; color: {color}; margin-bottom: 12px; }}
    .toc a  {{ display: block; color: #333; text-decoration: none; padding: 4px 0; font-size: 14px; }}
    .toc a:hover {{ color: {color}; }}
    .content {{ max-width: 900px; margin: 0 auto; padding: 0 20px 60px; }}
    .section {{ background: #fff; border-radius: 12px; padding: 32px; margin: 24px 0; box-shadow: 0 2px 8px rgba(0,0,0,0.07); }}
    h2.sec-heading {{ font-size: 22px; font-weight: 800; color: {color}; padding-bottom: 12px; border-bottom: 3px solid {color}; margin-bottom: 20px; }}
    h3.sub-heading {{ font-size: 15px; font-weight: 700; color: #333; margin: 18px 0 8px; }}
    p.body-text {{ font-size: 14px; line-height: 1.8; color: #444; margin-bottom: 12px; }}
    ul.bullets {{ margin-left: 20px; margin-bottom: 12px; }}
    ul.bullets li {{ font-size: 14px; line-height: 1.8; color: #444; margin-bottom: 4px; }}
    ol.numbered {{ margin-left: 20px; margin-bottom: 12px; }}
    ol.numbered li {{ font-size: 14px; line-height: 1.8; color: #444; margin-bottom: 4px; }}
    .info-box {{ background: #f0ebff; border-left: 4px solid {color}; padding: 12px 16px; border-radius: 0 8px 8px 0; margin: 12px 0; font-size: 13px; color: #444; }}
    .img-wrap {{ margin: 16px 0; text-align: center; }}
    .img-wrap img {{ max-width: 100%; border-radius: 10px; box-shadow: 0 4px 16px rgba(0,0,0,0.15); border: 1px solid #e0e0e0; }}
    .img-wrap figcaption {{ font-size: 12px; color: #888; margin-top: 6px; }}
    """

    sections_html = ""
    toc_items = ""
    for sec in doc["sections"]:
        h = sec["heading"]
        sid = re.sub(r"[^\w]", "_", h)
        toc_items += f'<a href="#{sid}">▸ {h}</a>\n'
        body = ""
        for item in sec.get("content", []):
            t = item[0]
            if t == "text":
                body += f'<p class="body-text">{item[1]}</p>\n'
            elif t == "heading3":
                body += f'<h3 class="sub-heading">{item[1]}</h3>\n'
            elif t == "bullets":
                lis = "".join(f"<li>{b}</li>" for b in item[1])
                body += f'<ul class="bullets">{lis}</ul>\n'
            elif t == "numbered":
                lis = "".join(f"<li>{b}</li>" for b in item[1])
                body += f'<ol class="numbered">{lis}</ol>\n'
            elif t == "info":
                body += f'<div class="info-box">ℹ️ {item[1]}</div>\n'
            elif t == "img":
                fname = item[1]
                cap   = item[2] if len(item) > 2 else fname
                b64   = img_to_b64(ss_dir / fname)
                if b64:
                    body += f'''<figure class="img-wrap">
  <img src="data:image/png;base64,{b64}" alt="{cap}">
  <figcaption>▲ {cap}</figcaption>
</figure>\n'''
        sections_html += f'''<div class="section" id="{sid}">
<h2 class="sec-heading">{h}</h2>
{body}
</div>\n'''

    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title>
<style>{css}</style>
</head>
<body>
<div class="cover">
  <h1>{title}</h1>
  <p>{sub}</p>
  <p style="margin-top:20px; opacity:0.6; font-size:13px;">KG스틸 당진생산지원팀 · {__import__('datetime').date.today()}</p>
</div>
<div class="toc">
  <h2>목 차</h2>
  {toc_items}
</div>
<div class="content">
{sections_html}
</div>
</body>
</html>"""


# ─────────────────────────────────────────
# Word 생성
# ─────────────────────────────────────────

def set_heading_style(para, level, color_hex):
    run = para.runs[0] if para.runs else para.add_run(para.text)
    run.font.color.rgb = RGBColor.from_string(color_hex.lstrip("#"))
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
    else:
        run.font.size = Pt(12)
        run.font.bold = True

def add_img_word(doc, img_path: Path, caption: str, width=Inches(5.5)):
    if not img_path.exists():
        return
    # 이미지 비율 유지하며 리사이즈
    with Image.open(img_path) as im:
        w, h = im.size
        ratio = h / w
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(img_path), width=width)
    cap = doc.add_paragraph(f"▲ {caption}")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

def render_word(doc_data: dict, ss_dir: Path, out_path: Path):
    color = doc_data["color"].lstrip("#")
    doc = Document()

    # 기본 스타일 설정
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)

    # 페이지 여백
    from docx.shared import Mm
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

    # 표지
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    t = doc.add_paragraph(doc_data["title"])
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].font.size = Pt(26)
    t.runs[0].font.bold = True
    t.runs[0].font.color.rgb = RGBColor.from_string(color)

    s = doc.add_paragraph(doc_data["subtitle"])
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.runs[0].font.size = Pt(13)
    s.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    import datetime
    d = doc.add_paragraph(f"KG스틸 당진생산지원팀 · {datetime.date.today()}")
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.runs[0].font.size = Pt(10)
    d.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.add_page_break()

    # 본문
    for sec in doc_data["sections"]:
        h = doc.add_heading(sec["heading"], level=1)
        h.runs[0].font.color.rgb = RGBColor.from_string(color)
        h.runs[0].font.size = Pt(16)

        for item in sec.get("content", []):
            t = item[0]
            if t == "text":
                p = doc.add_paragraph(item[1])
                p.runs[0].font.size = Pt(10.5)
            elif t == "heading3":
                h3 = doc.add_heading(item[1], level=2)
                h3.runs[0].font.size = Pt(12)
                h3.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            elif t == "bullets":
                for b in item[1]:
                    p = doc.add_paragraph(b, style="List Bullet")
                    p.runs[0].font.size = Pt(10.5)
            elif t == "numbered":
                for b in item[1]:
                    p = doc.add_paragraph(b, style="List Number")
                    p.runs[0].font.size = Pt(10.5)
            elif t == "info":
                p = doc.add_paragraph(f"ℹ️  {item[1]}")
                p.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x88)
                p.runs[0].font.italic = True
                p.runs[0].font.size = Pt(10)
            elif t == "img":
                fname = item[1]
                cap   = item[2] if len(item) > 2 else fname
                add_img_word(doc, ss_dir / fname, cap)

        doc.add_paragraph()

    doc.save(str(out_path))
    print(f"  Word 저장: {out_path.name}")


# ─────────────────────────────────────────
# 실행
# ─────────────────────────────────────────

if __name__ == "__main__":
    from PIL import Image  # pillow 필요

    print("=== 웹 매뉴얼 생성 ===")
    web_html = ROOT / "web_manual.html"
    web_docx = ROOT / "web_manual.docx"
    web_html.write_text(render_html(WEB_DOC, WEB_SS), encoding="utf-8")
    print(f"  HTML 저장: {web_html.name}")
    render_word(WEB_DOC, WEB_SS, web_docx)

    print("\n=== 앱 매뉴얼 생성 ===")
    app_html = ROOT / "app_manual.html"
    app_docx = ROOT / "app_manual.docx"
    app_html.write_text(render_html(APP_DOC, APP_SS), encoding="utf-8")
    print(f"  HTML 저장: {app_html.name}")
    render_word(APP_DOC, APP_SS, app_docx)

    print("\n완료!")
    print(f"  {web_html}")
    print(f"  {web_docx}")
    print(f"  {app_html}")
    print(f"  {app_docx}")
